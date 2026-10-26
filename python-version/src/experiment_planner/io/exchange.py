import csv
import copy
import hashlib
import json
import math
from contextlib import closing
from dataclasses import asdict, dataclass, field
from pathlib import Path

from experiment_planner.application.records import normalize_observations
from experiment_planner.metrics.uncertainty import Uncertainty
from experiment_planner.domain.errors import ValidationError
from experiment_planner.domain.template import parse_value
from experiment_planner.domain.reliability import batch_name, experiment_time, record_time
from experiment_planner.storage.project import encode
from experiment_planner.io.excel_guidance import (
    ExcelGuide, add_column_hints, add_value_validation, column_cells,
    has_hint_row, value_hint,
)


@dataclass
class ImportPreview:
    revision: int
    records: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    stop_row: int | None = None
    stop_fields: list = field(default_factory=list)
    mode: str = "append"
    changes: list = field(default_factory=list)

    @property
    def stop_message(self):
        if self.stop_row is None: return ""
        return f"第 {self.stop_row} 行的前 {len(self.stop_fields)} 个输入均为空，已作为结束标记；该行及后续行不导入。检查字段：" + "、".join(self.stop_fields)

    def commit(self, service):
        if self.errors: raise ValidationError("请先修复导入预览中的错误")
        if self.mode == "update": return service.update_records(self.records, self.revision)
        return service.import_records(self.records, self.revision)


def record_token(record):
    return hashlib.sha256(encode(record).encode("utf-8")).hexdigest()


def editable_record(record):
    """Canonical editable values; ignore import provenance and empty optional cells."""
    cells = {}
    for name, cell in record["observations"].items():
        value = {k: copy.deepcopy(cell.get(k)) for k in ("value", "unit", "note", "uncertainty")}
        value["uncertainty"] = asdict(Uncertainty(**(value["uncertainty"] or {})))
        value["note"] = value["note"] or ""
        if value["value"] is None and not value["note"] and (value["uncertainty"] or {}).get("kind", "unspecified") == "unspecified":
            continue
        cells[name] = value
    return {"conditions": record["actual"], "observations": cells,
            "status": record["status"], "note": record.get("note", ""),
            "field_notes": {n: v for n, v in record.get("field_notes", {}).items() if v},
            "experiment_batch": record.get("experiment_batch", ""), "performed_at": record_time(record)}


def record_changes(before, after, prefix=""):
    changes = []
    for key in sorted(set(before) | set(after)):
        old, new = before.get(key), after.get(key)
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(old, dict) and isinstance(new, dict):
            changes.extend(record_changes(old, new, path))
        elif old != new:
            changes.append({"field": path, "before": old, "after": new})
    return changes


def missing(value):
    return value is None or (isinstance(value, str) and not value.strip())


class TableHeaders(list):
    """A list of names with the original header's physical row."""
    def __init__(self, names, row_number):
        super().__init__(names)
        self.row_number = row_number


def checked_headers(values, row_number=1):
    headers = [str(value).strip() if value is not None else "" for value in values]
    while headers and not headers[-1]: headers.pop()
    if not headers or any(not h for h in headers) or len(headers) != len(set(headers)):
        raise ValidationError(f"导入表头为空或重复，请检查第 {row_number} 行")
    return TableHeaders(headers, row_number)


def table_rows(path):
    """Yield headers, then (physical starting row, cells), preserving empty rows."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source)
            headers = checked_headers(next(reader, []))
            yield headers
            while True:
                start = reader.line_num + 1
                values = next(reader, None)
                if values is None: break
                row = {name: values[i] if i < len(values) else None for i, name in enumerate(headers)}
                if any(not missing(value) for value in values[len(headers):]): row[None] = values[len(headers):]
                yield start, row
        return
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            values = next(rows, ())
            heading_row = 1
            if has_hint_row(values):
                heading_row = 2
                values = next(rows, ())
            headers = checked_headers(values, heading_row)
            yield headers
            for number, values in enumerate(rows, heading_row + 1):
                row = {name: values[i] if i < len(values) else None for i, name in enumerate(headers)}
                if any(not missing(value) for value in values[len(headers):]): row[None] = values[len(headers):]
                yield number, row
        finally: workbook.close()
        return
    raise ValidationError("仅支持 CSV / XLSX")


def read_headers(path):
    with closing(table_rows(path)) as rows: return next(rows)


def read_rows(path):
    with closing(table_rows(path)) as rows:
        next(rows)
        return [row for _, row in rows]


def preview_import(service, path, mapping=None, *, mode="append"):
    if mode not in ("append", "update"): raise ValidationError("请选择新增导入或批量修改")
    t = service.project.template
    preview = ImportPreview(service.project.revision, mode=mode)
    project_uid = service.project.db.execute("SELECT uuid FROM project WHERE id=1").fetchone()[0]
    seen = set()
    fields = t.parameters + t.measurements
    aliases = {f.get("label", f["name"]): f["name"] for f in fields}
    aliases.update({f["name"]: f["name"] for f in fields})
    suffixes = {"_不确定度": "_uncertainty", "_误差类型": "_uncertainty_kind", "_备注": "_note", "_单位": "_unit", "_分布": "_distribution", "_覆盖因子": "_coverage_factor", "_置信水平": "_confidence"}
    def mapped_name(column):
        if column in (mapping or {}): return mapping[column]
        for suffix, canonical in suffixes.items():
            if column.endswith(suffix) and column[:-len(suffix)] in aliases:
                return aliases[column[:-len(suffix)]] + canonical
        return aliases.get(column, column)
    count = max(1, math.ceil(len(t.parameters) * .3))
    first_inputs = [p["name"] for p in t.parameters[:count]]
    with closing(table_rows(path)) as rows:
        headers = next(rows)
        heading_row = headers.row_number
        mapped = {h: mapped_name(h) for h in headers}
        if len(set(mapped.values())) != len(mapped):
            preview.errors.append({"row": heading_row, "message": "多列映射到同一字段，请调整列对应关系"})
            return preview
        absent = [name for name in first_inputs if name not in mapped.values()]
        if mode == "update":
            required = ["experiment_id", "record_uid", "record_token", "project_uid", "status", "note",
                        "experiment_batch", "performed_at", *[f["name"] for f in t.parameters + t.measurements],
                        *[p["name"]+"_note" for p in t.parameters],
                        *[f["name"]+suffix for f in t.measurements for suffix in
                          ("_unit", "_note", "_uncertainty", "_uncertainty_kind", "_distribution", "_coverage_factor", "_confidence")]]
            absent += [name for name in required if name not in mapped.values() and name not in absent]
        if absent:
            preview.errors.append({"row": heading_row, "message": ("批量修改缺少必要列，请使用本程序最新导出的记录：" if mode == "update" else "用于结束判断的输入列未匹配：") + "、".join(absent)})
            return preview
        for row_number, raw in rows:
            row = {mapped[k]: value for k, value in raw.items() if k in mapped}
            if all(missing(row.get(name)) for name in first_inputs):
                preview.stop_row = row_number; preview.stop_fields = first_inputs
                break
            try:
                if None in raw: raise ValidationError("数据列数超过表头，请检查多余的分隔符或表头")
                old = None
                if mode == "update":
                    eid_value = parse_value(row.get("experiment_id"), {"name": "experiment_id", "value_type": "int"})
                    if eid_value is None or eid_value <= 0 or eid_value in seen:
                        raise ValidationError("实验编号无效或重复；批量修改每条实验只能出现一次")
                    seen.add(eid_value)
                    old = service.project.experiment(eid_value)
                    if (row.get("project_uid") != project_uid or (row.get("record_uid") or "") != old.get("record_uid", "")
                            or row.get("record_token") != record_token(old)):
                        raise ValidationError(f"实验 {eid_value} 身份或导出版本不匹配，请重新导出；请勿修改身份列或覆盖其他项目、已更正、已删除及编号复用的记录")
                inputs = {p["name"]: parse_value(row.get(p["name"]), p) for p in t.parameters}
                t.validate_conditions(inputs, execution=False)
                measurements = {}
                for f in t.measurements:
                    n = f["name"]
                    original = copy.deepcopy(old["observations"].get(n, {})) if old else {}
                    uncertainty = copy.deepcopy(original.get("uncertainty") or {})
                    if row.get(n+"_uncertainty_json"):
                        uncertainty = json.loads(row[n+"_uncertainty_json"])
                        if not isinstance(uncertainty, dict): raise ValidationError(f"{n} 的详细误差必须为 JSON 对象")
                    uncertainty.update({"kind": row.get(n+"_uncertainty_kind") or "unspecified", "distribution": row.get(n+"_distribution") or None})
                    for suffix in ("amount", "coverage_factor", "confidence"):
                        value = row.get(n+("_uncertainty" if suffix == "amount" else "_"+suffix))
                        if value not in ("", None): uncertainty[suffix] = float(value)
                        elif n+("_uncertainty" if suffix == "amount" else "_"+suffix) in row: uncertainty[suffix] = None
                    measurements[n] = {**original, "value": row.get(n), "unit": row.get(n+"_unit") or f.get("unit", "1"), "note": row.get(n+"_note") or "", "uncertainty": uncertainty}
                    if not old: measurements[n]["source_text"] = str(row.get(n, ""))
                measurements = normalize_observations(t, measurements)
                item = {"conditions": inputs, "observations": measurements, "external_id": row.get("external_id") or row.get("experiment_id"), "note": row.get("note") or "", "status": row.get("status") or "completed", "field_notes": {p["name"]: row.get(p["name"]+"_note") or "" for p in t.parameters},
                        "experiment_batch": batch_name(row.get("experiment_batch")), "performed_at": experiment_time(row.get("performed_at"))}
                if mode == "update":
                    for key in ("project_uid", "record_token"): item[key] = row[key]
                    item.update(experiment_id=eid_value, record_uid=row.get("record_uid") or "")
                checked = service.prepare_record(inputs, measurements, **{k: item[k] for k in ("status", "note", "field_notes", "experiment_batch", "performed_at")})
                if old:
                    changes = record_changes(editable_record(old), editable_record(checked))
                    preview.changes.append({"experiment_id": eid_value, "changes": changes})
                preview.records.append(item)
            except (ValueError, TypeError) as exc:
                preview.errors.append({"row": row_number, "message": str(exc)})
    return preview


def export_records(project, path):
    path = Path(path)
    if path.suffix.lower() == ".json":
        path.write_text(json.dumps(project.snapshot(), ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        return
    rows = []
    project_uid = project.db.execute("SELECT uuid FROM project WHERE id=1").fetchone()[0]
    for e in project.experiments():
        row = {"experiment_id": e["id"], "record_uid": e.get("record_uid", ""), "project_uid": project_uid,
               "record_token": record_token(e), "status": e["status"], "note": e["note"],
               "experiment_batch": e.get("experiment_batch", ""), "performed_at": record_time(e), **e["actual"]}
        for name, note in e.get("field_notes", {}).items(): row[name+"_note"] = note
        for n, c in e["observations"].items():
            u = c.get("uncertainty") or {}
            row.update({n: c["value"], n+"_unit": c["unit"], n+"_note": c["note"], n+"_uncertainty": u.get("amount"), n+"_uncertainty_kind": u.get("kind"), n+"_distribution": u.get("distribution"), n+"_coverage_factor": u.get("coverage_factor"), n+"_confidence": u.get("confidence"), n+"_uncertainty_json": json.dumps(u, ensure_ascii=False)})
        for n, c in e["derived"].items(): row.update({n: c["value"], n+"_status": c["status"]})
        row["suggested_json"] = json.dumps(e.get("suggested"), ensure_ascii=False)
        row["prediction_json"] = json.dumps(e.get("prediction"), ensure_ascii=False)
        rows.append(row)
    headers = list(dict.fromkeys(["experiment_id", "record_uid", "project_uid", "record_token", "status", "note", "experiment_batch", "performed_at",
        *(p["name"] for p in project.template.parameters), *(p["name"]+"_note" for p in project.template.parameters),
        *(f["name"] for f in project.template.measurements),
        *(f["name"]+suffix for f in project.template.measurements for suffix in
          ("_unit", "_note", "_uncertainty", "_uncertainty_kind", "_distribution", "_coverage_factor", "_confidence", "_uncertainty_json")),
        *(k for r in rows for k in r)]))
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        sheet = wb.active
        sheet.title = "实验记录"
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(k) for k in headers])
            # Notes and external text are literal text, never spreadsheet formula code.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str): cell.data_type = "s"
        guide = wb.create_sheet("批量修改说明")
        for values in [
            ["操作", "修改“实验记录”页后，在同一项目中导入并选择“批量修改已有实验”，核对逐项预览再保存。"],
            ["身份列", "experiment_id、record_uid、project_uid、record_token 请勿修改。记录更改、删除或编号复用后须重新导出。"],
            ["可修改", "输入条件、原始测量及误差、各备注、status、experiment_batch（自定义批次）、performed_at（实验日期）。"],
            ["留空", "已有测量留空表示清除；全部输入必须完整。保留列名，删除不需修改的整行可以只更新部分记录。"],
            ["实验日期", "填写 YYYY-MM-DD 或 ISO 日期时间；无时区按 UTC。完成/部分结果日期留空采用本次录入时间。"],
            ["批次可信度", "在软件“批次可信度 / 近期权重”中设置。未设置的批次默认 1。"],
            ["只读结果", "派生值由原始测量重算；prediction_json、suggested_json 仅供查看，修改它们不改变模型预测。"],
            ["误差详情", "常规误差列优先；uncertainty_json 保留相关来源及非对称界限等详细信息，通常无需改动。"],
        ]: guide.append(values)
        guide.column_dimensions["A"].width = 20; guide.column_dimensions["B"].width = 110
        sheet.freeze_panes = "E2"; sheet.auto_filter.ref = sheet.dimensions
        wb.active = 0
        wb.save(path)
        wb.close()
    else: raise ValidationError("导出支持 CSV / XLSX / JSON")


def write_table(path, headers, rows, title):
    """Write numeric values at full precision, and spreadsheet text literally."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=headers)
            writer.writeheader(); writer.writerows(rows)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active; sheet.title = title
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key) for key in headers])
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str): cell.data_type = "s"
        sheet.freeze_panes = "A2"
        workbook.save(path); workbook.close()
    else:
        raise ValidationError("请选择 CSV 或 Excel (.xlsx)")


def export_importance(report, template, path):
    labels = {f["name"]: f.get("label", f["name"]) for f in template.fields}
    rows = []
    for ranked in report["ranking"]:
        for response in report["responses"]:
            item = next(row for row in response["parameters"] if row["parameter"] == ranked["parameter"])
            rows.append({"参数": ranked["parameter"], "参数名称": ranked["label"], "综合重要度": ranked["score"],
                         "响应": response["response"], "响应名称": labels.get(response["response"], response["response"]),
                         "响应重要度": item["score"], "敏感度原值": item["sensitivity"], "说明": item["reason"],
                         "训练条数": response["count"], "分析条数": response["analyzed_count"],
                         "数据版本": report["revision"], "模型": report["model"], "置换次数": report["repeats"],
                         "解释": " ".join(report["notices"][:3])})
    headers = list(rows[0]) if rows else ["参数", "综合重要度"]
    write_table(path, headers, rows, "实验条件重要度")


def export_import_template(template, path):
    """An empty input file with visible Excel hints and cell validation."""
    headers = ["external_id", "status", "note", "experiment_batch", "performed_at"]
    for field in template.parameters:
        headers.extend([field["name"], field["name"] + "_note"])
    for field in template.measurements:
        name = field["name"]
        headers.extend([name, *[name + suffix for suffix in
            ("_unit", "_uncertainty", "_uncertainty_kind", "_distribution",
             "_coverage_factor", "_confidence", "_note")]])
    checked_headers(headers)
    path = Path(path)
    write_table(path, headers, [], "导入数据")
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path)
        _guide_import_sheet(wb, template, headers)
        help_sheet = wb.create_sheet("填写说明")
        help_sheet.append(["模板", template.data.get("name", ""), "版本", template.data["template_version"]])
        help_sheet.append(["提示", "只在“导入数据”填写记录；不要在数据中间留空行。每次复测填写不同 external_id。"])
        help_sheet.append(["行与限定", "第 1 行是填写提示，第 2 行是列名，第 3 行起填写。下拉及数值校验覆盖 10000 行；粘贴可能绕过 Excel 校验，最终以软件预览为准。"])
        help_sheet.append(["必填", "每条记录所有输入参数必填，固定输入也要填。测量与误差可暂缺；记录可保存，完整性决定能否用于模型。"])
        help_sheet.append(["状态", "留空为 completed；可填 completed / partial / pending / running / failed / cancelled。"])
        help_sheet.append(["误差类型", "std / sem / bounds / interval / unspecified；只填 ± 数字不代表已知标准误差。"])
        help_sheet.append(["列名", "显示名称", "单位", "类型", "范围 / 取值", "备注"])
        for field in template.parameters + template.measurements:
            help_sheet.append([field["name"], field.get("label", field["name"]), field.get("unit", "1"),
                field.get("value_type", "float"), json.dumps(field.get("bounds", field.get("values", field.get("fixed_value"))), ensure_ascii=False), field.get("note", "")])
        for cells in help_sheet:
            for cell in cells:
                if isinstance(cell.value, str): cell.data_type = "s"
        wb.active = 0; wb.save(path); wb.close()


def _guide_import_sheet(workbook, template, headers):
    sheet = workbook.active
    guide = ExcelGuide(workbook)
    hints = {
        "external_id": "选填：外部实验编号。复测使用不同编号；已有编号用于重复导入检查。",
        "status": "选填：从下拉选择状态；留空默认为 completed（已完成）。",
        "note": "选填：整条实验备注。",
        "experiment_batch": "选填：自定义实验批次名称。同名批次共用在软件设置的可信度，未设置为 1。",
        "performed_at": "选填：实际实验日期 YYYY-MM-DD 或 ISO 日期时间；完成/部分结果留空采用录入时间，无时区按 UTC。",
    }
    choices = {"status": ["completed", "partial", "pending", "running", "failed", "cancelled"]}
    numeric = {}
    for field in template.parameters:
        name = field["name"]
        hints[name] = value_hint(field, required=True)
        hints[name + "_note"] = "选填：该输入的备注。实际数值不要求落在推荐执行网格上。"
    for field in template.measurements:
        name = field["name"]
        hints[name] = value_hint(field, limits=False) + "未测可留空。"
        hints.update({
            name + "_unit": f"选填：单位只能为 {field.get('unit', '1')}，留空使用模板单位。更换单位请先修改模板。",
            name + "_uncertainty": "选填：不确定度数值 ≥ 0，使用本测量单位。有误差类型时按类型填写；未测量时留空。",
            name + "_uncertainty_kind": "选填：std 标准差、sem 均值标准误、bounds 对称界限、interval 区间、unspecified 未知。需配合不确定度数值解释大小。",
            name + "_distribution": "选填：normal 正态、uniform 均匀。bounds + uniform 才转换为标准误差；interval + normal 需覆盖因子或置信水平。",
            name + "_coverage_factor": "选填：数值 > 0。interval + normal 若需换算标准误差，与置信水平至少填一个；否则误差按未知处理。",
            name + "_confidence": "选填：0 < 数值 < 1，例如 0.95。interval + normal 若需换算标准误差，与覆盖因子至少填一个。",
            name + "_note": "选填：该测量备注。",
        })
        choices[name + "_unit"] = [field.get("unit", "1")]
        if field.get("value_type", "float") in ("float", "int"):
            choices[name + "_uncertainty_kind"] = ["unspecified", "std", "sem", "bounds", "interval"]
            choices[name + "_distribution"] = ["normal", "uniform"]
            numeric[name + "_uncertainty"] = {"lower": 0}
            numeric[name + "_coverage_factor"] = {"lower": 0, "strict_lower": True}
            numeric[name + "_confidence"] = {"lower": 0, "upper": 1, "strict_lower": True, "strict_upper": True}
        else:
            choices[name + "_uncertainty_kind"] = ["unspecified"]
            for suffix in ("_uncertainty", "_uncertainty_kind", "_distribution", "_coverage_factor", "_confidence"):
                hints[name + suffix] = "非数值测量：留空，不填写数值误差；误差类型只能 unspecified（未知）。"
    add_column_hints(sheet, [hints[key] for key in headers])
    for index, key in enumerate(headers, 1):
        if key in choices: guide.dropdown(sheet, column_cells(index), choices[key], hints[key])
        if key in numeric: guide.number(sheet, column_cells(index), hints[key], **numeric[key])
    for field in template.parameters + template.measurements:
        index = headers.index(field["name"]) + 1
        add_value_validation(guide, sheet, index, field, hints[field["name"]], limits=field in template.parameters)


def export_analysis(report, template, path):
    """Export exactly the cached prediction and plotted samples; do not refit."""
    rows = []
    for source, data in (("预测", report), ("趋势图", report.get("plot", {}))):
        for index, (condition, prediction) in enumerate(zip(data.get("conditions", []), data.get("predictions", [])), 1):
            variants = [("输入确定", prediction)]
            if "nominal_prediction" in prediction:
                variants = [("输入确定", prediction["nominal_prediction"]), ("含输入执行波动", prediction)]
            for interpretation, variant in variants:
                for name, item in variant.items():
                    if not isinstance(item, dict) or "quantiles" not in item: continue
                    q = item.get("quantiles") or [None, None, None]
                    field = next((f for f in template.fields if f["name"] == name), {})
                    rows.append({"来源": source, "点序号": index, "数据版本": report["revision"],
                        "模板版本": report["template_version"], "模型": report["model"], "随机种子": variant.get("seed"),
                        **condition, "指标": name, "指标名称": field.get("label", name), "单位": field.get("unit", "1"),
                        "后验均值": item.get("mean"), "下限2.5%": q[0], "中位数50%": q[1], "上限97.5%": q[2],
                        "区间类型": item.get("interval_kind"), "预测可行概率": variant.get("feasibility_probability"),
                        "有限样本数": item.get("finite_sample_count"), "非有限样本数": item.get("nonfinite_sample_count"),
                        "计算口径": interpretation,
                        "输入不确定度配置": json.dumps(prediction["input_uncertainty"], ensure_ascii=False, allow_nan=False) if prediction.get("input_uncertainty") else "未启用",
                        "说明": item.get("moment_notice", "")})
    if not rows: raise ValidationError("请先完成预测或趋势图计算")
    headers = list(rows[0])
    write_table(path, headers, rows, "预测与趋势数据")
