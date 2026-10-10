"""Visible instructions and Excel validation shared by the two template formats."""
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


HINT_PREFIX = "【填写提示】"
FIRST_DATA_ROW = 3
LAST_DATA_ROW = 10002


def has_hint_row(values):
    return bool(values) and str(values[0] or "").startswith(HINT_PREFIX)


def header_row(sheet):
    return 2 if has_hint_row([sheet.cell(1, 1).value]) else 1


def add_column_hints(sheet, hints):
    """Put human guidance above unchanged machine-readable column names."""
    sheet.insert_rows(1)
    for index, hint in enumerate(hints, 1):
        cell = sheet.cell(1, index, (HINT_PREFIX if index == 1 else "") + hint)
        cell.data_type = "s"
        cell.font = Font(size=10, color="B91C1C" if "必填" in hint else "78350F")
        cell.fill = PatternFill("solid", fgColor="FEF3C7")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        heading = sheet.cell(2, index)
        heading.font = Font(bold=True, color="FFFFFF")
        heading.fill = PatternFill("solid", fgColor="2563EB")
        heading.alignment = Alignment(wrap_text=True, vertical="center")
        heading.comment = Comment(hint, "实验规划助手")
        sheet.column_dimensions[get_column_letter(index)].width = 24
    sheet.row_dimensions[1].height = 110
    sheet.row_dimensions[2].height = 32
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(hints))}{max(2, sheet.max_row)}"


class ExcelGuide:
    def __init__(self, workbook):
        self.workbook = workbook
        self._options = {}

    def _validation(self, sheet, cells, hint, **kwargs):
        validation = DataValidation(
            allow_blank=True, showErrorMessage=True, showInputMessage=True,
            errorStyle="stop", errorTitle="填写内容不符合要求",
            error="请按列名上方说明填写。复制粘贴后仍需通过软件导入校验。",
            promptTitle="填写说明", prompt=hint[:255], **kwargs,
        )
        sheet.add_data_validation(validation)
        validation.add(cells)
        return validation

    def dropdown(self, sheet, cells, options, hint):
        # Named ranges avoid Excel's 255-character inline-list limit and keep
        # commas, non-ASCII text and strings beginning '=' literal.
        options = tuple(dict.fromkeys(str(v) for v in options))
        if not options: return
        if options not in self._options:
            name = f"_planner_options_{len(self._options) + 1}"
            source = self.workbook["_填写选项"] if "_填写选项" in self.workbook else self.workbook.create_sheet("_填写选项")
            source.sheet_state = "hidden"
            column = len(self._options) + 1
            for index, option in enumerate(options, 1):
                cell = source.cell(index, column, option); cell.data_type = "s"
            letter = get_column_letter(column)
            self.workbook.defined_names.add(DefinedName(
                name, attr_text=f"'_填写选项'!${letter}$1:${letter}${len(options)}"))
            self._options[options] = name
        self._validation(sheet, cells, hint, type="list",
                         formula1=f"={self._options[options]}", showDropDown=False)

    def number(self, sheet, cells, hint, *, integer=False, lower=None, upper=None,
               strict_lower=False, strict_upper=False):
        first = cells.split(":")[0]
        conditions = [f"ISNUMBER({first})"]
        if integer: conditions.append(f"{first}=INT({first})")
        if lower is not None: conditions.append(f"{first}{'>' if strict_lower else '>='}{lower}")
        if upper is not None: conditions.append(f"{first}{'<' if strict_upper else '<='}{upper}")
        formula = f'OR({first}="",AND({",".join(conditions)}))'
        self._validation(sheet, cells, hint, type="custom", formula1=formula)


def column_cells(index):
    letter = get_column_letter(index)
    return f"{letter}{FIRST_DATA_ROW}:{letter}{LAST_DATA_ROW}"


def value_hint(field, required=False, *, limits=True):
    kind = field.get("value_type", "float")
    label = field.get("label", field["name"])
    hint = f"{'必填' if required else '选填'}：{label}；单位 {field.get('unit', '1')}。"
    if (limits or kind == "category") and "fixed_value" in field: hint += f"只能填 {field['fixed_value']}。"
    elif (limits or kind == "category") and "values" in field: hint += "从下拉选择限定值。"
    elif kind == "bool": hint += "从下拉选择 true / false（是 / 否）。"
    elif kind in ("float", "int"):
        hint += "整数。" if kind == "int" else "数值。"
        if limits and "bounds" in field:
            hint += f"允许范围：{field['bounds'][0]} ≤ 值 ≤ {field['bounds'][1]}（含边界）。"
    else: hint += "填写文字。"
    if field.get("note"): hint += str(field["note"])
    return hint


def add_value_validation(guide, sheet, index, field, hint, *, limits=True):
    cells = column_cells(index)
    kind = field.get("value_type", "float")
    if (limits or kind == "category") and "fixed_value" in field:
        values = [field["fixed_value"]]
    elif (limits or kind == "category") and "values" in field:
        values = field["values"]
    elif kind == "bool": values = ["true", "false", "是", "否"]
    else: values = None
    if values is not None:
        if kind == "bool": values = [str(v).lower() if type(v) is bool else v for v in values]
        guide.dropdown(sheet, cells, values, hint)
    elif kind in ("float", "int"):
        bounds = field.get("bounds", [None, None]) if limits else [None, None]
        guide.number(sheet, cells, hint, integer=kind == "int", lower=bounds[0], upper=bounds[1])
