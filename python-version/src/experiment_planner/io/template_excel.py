"""Editable Chinese Excel template definitions (distinct from experiment rows)."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.comments import Comment

from experiment_planner.application.template_editing import blank_template, normalize_draft
from experiment_planner.domain.errors import ValidationError
from experiment_planner.io.excel_guidance import ExcelGuide, add_column_hints, boolean_text, column_cells, header_row


FIELD_COLUMNS = [
    ("内部名称", "name"), ("显示名称", "label"), ("类型", "value_type"), ("单位", "unit"),
    ("说明", "note"), ("下限", "lower"), ("上限", "upper"), ("离散值", "values"),
    ("固定值", "fixed_value"), ("作为模型输入", "use_as_model_input"),
    ("作为响应", "is_response"), ("直接建模", "model_base"), ("测量时点", "measured_at"),
    ("计算公式", "expression"), ("显示小数位数", "display_digits"),
    ("执行模式", "execution_mode"), ("执行小数位数", "execution_digits"),
    ("步长", "execution_step"), ("网格起点", "execution_origin"), ("其他属性", "extra"),
]
TABLES = {
    "输入参数": ("parameters", FIELD_COLUMNS),
    "原始测量": ("measurements", FIELD_COLUMNS),
    "派生指标与公式": ("derived_metrics", FIELD_COLUMNS),
    "优化目标": ("objectives", [("响应或指标", "metric"), ("方向", "direction"), ("变换", "transform"), ("目标值", "target"), ("其他属性", "extra")]),
    "先验关系": ("knowledge_priors", [("启用", "enabled"), ("输入参数", "input"), ("基础响应", "response"), ("关系", "relation"), ("系数", "coefficient"), ("强度", "strength"), ("来源", "source"), ("其他属性", "extra")]),
    "输入约束": ("input_constraints", [("系数", "coefficients"), ("比较", "op"), ("阈值", "rhs"), ("其他属性", "extra")]),
    "输出约束": ("outcome_constraints", [("响应或指标", "metric"), ("比较", "op"), ("阈值", "bound"), ("其他属性", "extra")]),
}
ALIASES = {
    "数值": "float", "整数": "int", "类别": "category", "是否": "bool", "文字": "text",
    "实验前": "before", "实验后": "after", "未指定": "unknown",
    "不限制": "none", "小数位数": "digits", "指定步长": "step",
    "尽量大": "maximize", "尽量小": "minimize", "原值": "identity", "绝对值": "abs", "距指定值的距离": "absolute_distance",
    "多目标折中": "pareto", "单目标": "single", "加权目标": "weighted",
    "移除比值目标": "remove", "稳定化比值": "stabilized", "有效区域内比值": "valid_region",
}
BOOL_KEYS = {"use_as_model_input", "is_response", "model_base", "enabled"}
INT_KEYS = {"display_digits", "execution_digits", "schema_version", "template_version"}
NUMBER_KEYS = {"lower", "upper", "execution_step", "execution_origin", "target", "coefficient", "strength", "rhs", "bound"}

CHOICES = {
    "value_type": ["float", "int", "category", "text", "数值", "整数", "类别", "是否", "文字"],
    "measured_at": ["before", "after", "unknown", "实验前", "实验后", "未指定"],
    "execution_mode": ["none", "digits", "step", "不限制", "小数位数", "指定步长"],
    "direction": ["maximize", "minimize", "尽量大", "尽量小"],
    "transform": ["identity", "abs", "absolute_distance", "原值", "绝对值", "距指定值的距离"],
    "relation": ["independent", "linear", "proportional", "monotonic", "weak"],
    "op": ["<=", ">="],
    **{key: ["是", "否"] for key in BOOL_KEYS},
}


def _column_hint(section, key):
    hints = {
        "name": "必填：英文字母开头，仅字母、数字、下划线；所有字段表不得重名。",
        "label": "选填：显示名称；留空使用内部名称。", "unit": "选填：自由输入单位；留空为 1（无量纲），不自动换算自定义单位。",
        "value_type": "选填：下拉选数值 / 整数 / 类别 / 是否 / 文字；是否类型的值填是或否。留空为 float（数值）。",
        "note": "选填：字段说明。", "lower": "条件必填：用范围定义数值输入时，下限和上限都要填，且下限 < 上限。范围、离散值、固定值三选一。",
        "upper": "条件必填：范围输入的上限；必须大于下限。无范围时留空。",
        "values": '条件必填：类别输入需要离散值或固定值；填写 JSON 数组，例如 ["甲","乙"] 或 [1,2,3]。',
        "fixed_value": "选填：固定输入只填该值，不再填范围或离散值；按字段类型填写。",
        "use_as_model_input": "选填：是 或 否。文字不能作为数值模型输入；事后或未知测量不能作为未来输入。",
        "is_response": "选填：作为响应？原始测量/派生指标的数值字段手写留空默认是；文字/类别不作数值响应。",
        "model_base": "选填：直接建模？原始测量手写留空默认是；派生指标留空默认否。至少一个响应需直接建模。",
        "measured_at": "选填：测量时点 before/实验前、after/实验后、unknown/未指定。手写原始测量留空默认 after。",
        "expression": "必填（派生指标）：用内部名称写公式，例如 quality * 2。其他字段留空。",
        "display_digits": "选填：整数范围 -12 ≤ 位数 ≤ 12；空白为原始精度，仅改变显示。",
        "execution_mode": "选填：none 不限制、digits 小数位数、step 步长。仅数值输入使用；空白默认 none。",
        "execution_digits": "条件必填：digits 模式填整数，-12 ≤ 位数 ≤ 12；不用时留空。",
        "execution_step": "条件必填：step 模式的步长，数值 > 0。",
        "execution_origin": "选填：step 网格起点，任意有限数值；留空为 0。",
        "extra": "选填：扩展属性 JSON 对象；不清楚可留空。",
        "metric": "必填：已有响应或派生指标的内部名称。", "direction": "选填：maximize 尽量大、minimize 尽量小；留空为 maximize。",
        "transform": "选填：identity 原值、abs 绝对值、absolute_distance 距指定值的距离；留空为 identity。",
        "target": "条件必填：变换为 absolute_distance 时的目标值；任意有限数值。",
        "enabled": "选填：是 启用、否 不启用；留空默认启用。",
        "input": "启用时必填：数值模型输入的内部名称。", "response": "启用时必填：直接建模基础响应的内部名称。",
        "relation": "启用时必填：independent 独立、linear 线性、proportional 比例、monotonic 单调、weak 弱先验。",
        "coefficient": "条件必填：启用非 independent/weak 关系时填有符号系数；任意有限数值，单位为响应单位/输入单位。",
        "strength": "条件必填：启用非 independent 关系时的强度，数值 > 0。",
        "source": "启用时必填：经验或文献依据。", "coefficients": '必填：非空 JSON 系数组，例如 {"speed":1}；内部名称须是数值输入。',
        "op": "必填：从下拉选 <=（不超过）或 >=（不低于）。",
        "rhs": "必填：输入约束阈值，任意有限数值。", "bound": "必填：输出约束阈值，任意有限数值。",
    }
    if section != "derived_metrics" and key == "expression": return "留空：本区域不填写派生公式。"
    if section != "parameters" and key in {"lower", "upper", "values", "fixed_value", "execution_mode", "execution_digits", "execution_step", "execution_origin"}:
        if key == "values": return '条件必填：类别测量填写允许值的 JSON 数组，例如 ["甲","乙"]。'
        return "选填：非输入字段通常留空。范围和执行网格仅用于可控输入；下限/上限若填写须成对。"
    return hints[key]


SETTING_CHOICES = {
    "input_uncertainty.samples": [8, 16, 32],
    "optimization.mode": ["pareto", "single", "weighted"],
    "optimization.default_model_preset": ["gp_rbf_v1", "gp_matern25_v1", "bayesian_linear_v1"],
    "ratio_policy.mode": ["remove", "stabilized", "valid_region", "require_configuration"],
    "batch_defaults.search_mode": ["full_space", "single", "double", "cross"],
}
SETTING_NUMBERS = {
    "schema_version": (2, 2, True), "template_version": (1, None, True),
    "ratio_policy.epsilon_nm": (0, None, False),
    "batch_defaults.total_count": (1, 100, True), "batch_defaults.exploration_count": (0, 100, True),
    "batch_defaults.pool_size": (1, 10000, True), "batch_defaults.seed": (0, 2147483647, True),
    "batch_defaults.exploration_strength": (0, 100, False),
}


def _setting_hint(key):
    notes = {
        "input_uncertainty.enabled": "可选总开关，默认否。开启后对所选连续输入考虑未来执行波动，历史输入仍按记录值训练。",
        "input_uncertainty.samples": "每个条件的输入扰动采样数：8、16 或 32。分布按模板范围与线性约束条件化，数值越大计算通常越慢。",
        "name": "模板名称，可自由修改。", "schema_version": "必填：格式版本，固定为整数 2。",
        "template_version": "必填：模板版本，整数 ≥ 1。",
        "optimization.mode": "必填：pareto 多目标折中；single 恰好一个有效目标；weighted 配合 weights/scales 的 JSON 数组。",
        "optimization.default_model_preset": "默认模型：gp_rbf_v1 RBF、gp_matern25_v1 Matérn 2.5、bayesian_linear_v1 贝叶斯线性。",
        "optimization.weights": "条件必填：weighted 模式每个有效目标一个非负权重，不能全部为零；JSON 数组，顺序对应有效目标。",
        "optimization.scales": "条件必填：weighted 模式每个有效目标一个正数尺度（>0）；JSON 数组，顺序对应有效目标。",
        "ratio_policy.mode": "remove 移除指定比值目标；stabilized 稳定化；valid_region 有效区域；require_configuration 待配置。",
        "ratio_policy.epsilon_nm": "条件必填：stabilized/valid_region 时数值 > 0；按实际分母单位口径填写。",
        "modeling.base_responses": '说明性记录，当前实际基础响应由字段“作为响应 + 直接建模”决定。新实验可删除本行，或改为新内部名称的 JSON 数组，例如 ["quality"]。',
        "context_assumptions": "说明性场景假设，不自动排除模型输入或改变误差计算。全新实验可删除本行或填 []，有相关假设才重新填写。详见 回答.md。",
    }
    if key in notes: return notes[key]
    if key in SETTING_NUMBERS:
        lower, upper, integer = SETTING_NUMBERS[key]
        return f"选填：{'整数' if integer else '数值'}范围 {lower} ≤ 值" + (f" ≤ {upper}。" if upper is not None else "。")
    if key in SETTING_CHOICES: return "选填：从下拉选择固定选项。"
    return "选填：列表/对象写 JSON；字段引用须改为本实验的内部名称。"


def _guide_definition_sheet(guide, sheet, section, columns):
    hints = [_column_hint(section, key) for _, key in columns]
    add_column_hints(sheet, hints)
    for index, (_, key) in enumerate(columns, 1):
        cells = column_cells(index); hint = hints[index - 1]
        if key in CHOICES: guide.dropdown(sheet, cells, CHOICES[key], hint)
        elif key in {"display_digits", "execution_digits"}: guide.number(sheet, cells, hint, integer=True, lower=-12, upper=12)
        elif key in NUMBER_KEYS:
            positive = key in {"execution_step", "strength"}
            guide.number(sheet, cells, hint, lower=0 if positive else None, strict_lower=positive)


def _cell_value(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, (dict, list)) else value


def _append(sheet, values):
    sheet.append([_cell_value(value) for value in values])
    # Store strings literally, including expressions and user notes beginning '='.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str): cell.data_type = "s"


def _flatten(data, prefix=""):
    for key, value in data.items():
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict) and value:
            yield from _flatten(value, path)
        else:
            yield path, value


def write_template_excel(data, path):
    data = normalize_draft(data)
    workbook = Workbook()
    guide = ExcelGuide(workbook)
    help_sheet = workbook.active; help_sheet.title = "填写说明"
    for text in (
        "本文件定义实验模板，不是实验记录导入表。可导入“新建 / 复制 / 修改模板”继续编辑。",
        "保留中文表头；每行一个字段，内部名称以英文字母开头，可用数字和下划线，所有表中不得重名。",
        "各表第 1 行是填写提示，第 2 行是列名，第 3 行起填数据；不要修改表头。固定选项有下拉，数值限定见列名上方。",
        "必填与条件必填在提示行标出；限定覆盖 10000 行。复制粘贴可能绕过 Excel 校验，导入和最终保存仍由软件校验。",
        "类型可选数值/整数/类别/是否/文字；是否的填写值选是或否。单位可自由填写，无量纲填 1。",
        "数值输入填写下限/上限，或离散值，或固定值；这三种方式选一种。是否输入可全部留空。",
        '离散值可写 JSON 数组（例如 [1,2,3] 或 ["甲","乙"]），也可用中英文逗号分隔。',
        "文字输入不进入数值模型；可自由填写，推荐时需基准文字或固定值/离散值。",
        "开关列只需选 是/否；是否类型的固定值填是或否，离散值可写 [\"是\",\"否\"]。文字/类别测量的作为响应填否。",
        "原始测量作为响应、直接建模默认是；派生指标作为响应默认是、直接建模默认否。",
        "派生公式使用内部名称，例如 result * 2；可写 =result * 2，导入只解析文字，不执行 Excel 公式。",
        "至少一个输入、一个直接建模的响应、一个有效目标。目标方向填 maximize/minimize 或 尽量大/尽量小。",
        "模板设置每行一个属性路径，例如 name、optimization.mode（pareto/single/weighted）、ratio_policy.mode。",
        "模板设置的列表/对象填写 JSON；加权模式用 optimization.weights/scales，顺序对应有效目标。",
        "其他属性可留空；用于保留扩展 JSON 配置。输入约束系数示例：{\"speed\":1}。",
        "留空属性表示待填写，可先导入编辑；最终保存时校验类型、范围、公式量纲和优化目标。",
        "显示小数位数仅改变显示；执行模式 digits 配合执行小数位数，step 配合步长与网格起点。",
    ): _append(help_sheet, [text])
    for title, (section, columns) in TABLES.items():
        sheet = workbook.create_sheet(title)
        _append(sheet, [title for title, _ in columns])
        for field in data.get(section, []):
            row = dict(field)
            if columns is FIELD_COLUMNS:
                if section != "parameters":
                    # Export effective roles explicitly: handwritten worksheets
                    # have helpful defaults, but export/import must not activate
                    # an originally inactive field or stop direct modeling.
                    row.setdefault("is_response", False)
                    row.setdefault("model_base", True)
                    if section == "measurements": row.setdefault("measured_at", "unknown")
                if "bounds" in row: row["lower"], row["upper"] = row.pop("bounds")
                rule = row.pop("execution_rounding", {})
                for key in ("mode", "digits", "step", "origin"):
                    if key in rule: row[f"execution_{key}"] = rule[key]
                extra_rule = {k: v for k, v in rule.items() if k not in ("mode", "digits", "step", "origin")}
            else: extra_rule = {}
            extra = {k: v for k, v in row.items() if k not in {key for _, key in columns} or v is None}
            if extra_rule: extra["execution_rounding"] = extra_rule
            row["extra"] = extra or None
            if row.get("value_type") == "bool":
                row["value_type"] = "是否"
                if "fixed_value" in row: row["fixed_value"] = boolean_text(row["fixed_value"])
                if "values" in row: row["values"] = [boolean_text(v) for v in row["values"]]
            _append(sheet, [boolean_text(row.get(key)) if key in BOOL_KEYS else row.get(key) for _, key in columns])
        _guide_definition_sheet(guide, sheet, section, columns)
    settings = workbook.create_sheet("模板设置")
    _append(settings, ["属性", "值", "类型（可留空）", "填写说明"])
    for key, value in _flatten({k: v for k, v in data.items() if k not in {s for s, _ in TABLES.values()}}):
        # JSON null distinguishes a stored null from a new empty cell.
        _append(settings, [key, "null" if value is None else boolean_text(value),
                           "null" if value is None else "是否" if type(value) is bool else type(value).__name__, _setting_hint(key)])
    add_column_hints(settings, ["必填：属性路径，例如 name、optimization.mode。全新实验删除无关的旧元数据行。",
        "必填或选填见本行右侧说明；存在固定选项时从下拉选择。列表/对象填写 JSON。",
        "选填：str 文字 / int 整数 / float 数值 / 是否 / list 列表 / dict 对象 / null 空值。",
        "仅供阅读：本行的必填条件、范围、作用；本列不导入模板配置。"])
    guide.dropdown(settings, column_cells(3), ["str", "int", "float", "是否", "list", "dict", "null"], "属性值的类型；留空按 JSON 或文字解析。")
    for row in range(3, settings.max_row + 1):
        key, _, kind, hint = [settings.cell(row, column).value for column in range(1, 5)]
        cell = settings.cell(row, 2); cell.comment = Comment(hint, "实验规划助手")
        cells = cell.coordinate
        if key in SETTING_CHOICES: guide.dropdown(settings, cells, SETTING_CHOICES[key], hint)
        elif key in SETTING_NUMBERS:
            lo, hi, integer = SETTING_NUMBERS[key]
            guide.number(settings, cells, hint, integer=integer, lower=lo, upper=hi, strict_lower=key == "ratio_policy.epsilon_nm")
        elif kind in ("bool", "是否"): guide.dropdown(settings, cells, ["是", "否"], hint)
    settings.column_dimensions["A"].width = 42
    settings.column_dimensions["B"].width = 60
    settings.column_dimensions["D"].width = 80
    for sheet in workbook:
        if sheet.sheet_state == "hidden": continue
        if sheet is help_sheet:
            sheet.column_dimensions["A"].width = 110
        elif sheet is settings:
            for cells in sheet.iter_rows(min_row=3):
                cells[3].alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[cells[0].row].height = 55
    workbook.save(path)
    workbook.close()


def _parse(key, value):
    if value is None or (isinstance(value, str) and not value.strip()): return None
    if isinstance(value, str): value = value.strip()
    if key in BOOL_KEYS:
        if type(value) is bool: return value
        mapping = {"是": True, "否": False, "true": True, "false": False, "1": True, "0": False}
        if str(value).lower() not in mapping: raise ValidationError("开关需填写是/否或 true/false")
        return mapping[str(value).lower()]
    if key in INT_KEYS:
        number = float(value)
        if not number.is_integer(): raise ValidationError("小数位数和版本需为整数")
        return int(number)
    if key in NUMBER_KEYS: return float(value)
    if key in ("extra", "coefficients"):
        result = json.loads(value) if isinstance(value, str) else value
        if not isinstance(result, dict): raise ValidationError("需要 JSON 属性组")
        return result
    if key == "values":
        if isinstance(value, str) and value.startswith("["): return json.loads(value)
        return [v.strip() for v in str(value).replace("，", ",").split(",") if v.strip()]
    if key in ("value_type", "measured_at", "execution_mode", "direction", "transform"):
        return ALIASES.get(value, value)
    return value


def read_template_excel(path):
    workbook = load_workbook(path, data_only=False)
    try:
        if not any(title in workbook.sheetnames for title in TABLES):
            raise ValidationError("不是模板定义 Excel：需要“输入参数”“原始测量”等工作表，可先导出模板作为填写格式")
        data = blank_template()
        for title, (section, columns) in TABLES.items():
            if title not in workbook: continue
            sheet = workbook[title]
            heading_row = header_row(sheet)
            headers = {cell.value: index for index, cell in enumerate(sheet[heading_row]) if cell.value is not None}
            expected = {label for label, _ in columns} | {key for _, key in columns}
            unknown = set(headers) - expected
            if unknown: raise ValidationError(f"{title}：未知表头 {', '.join(map(str, unknown))}")
            required = "内部名称" if columns is FIELD_COLUMNS else "响应或指标" if section in ("objectives", "outcome_constraints") else "输入参数" if section == "knowledge_priors" else "系数"
            if required not in headers and dict(columns)[required] not in headers:
                raise ValidationError(f"{title}：缺少表头 {required}")
            for row_number, cells in enumerate(sheet.iter_rows(min_row=heading_row + 1), heading_row + 1):
                if all(c.value is None or str(c.value).strip() == "" for c in cells): continue
                try:
                    row = {}
                    for label, key in columns:
                        index = headers.get(label, headers.get(key))
                        if index is None: continue
                        cell = cells[index]
                        value = cell.value
                        if cell.data_type == "f" and key != "expression":
                            raise ValidationError(f"{label} 请填写实际值，不能使用 Excel 单元格公式")
                        if key == "expression" and isinstance(value, str): value = value.removeprefix("=")
                        value = _parse(key, value)
                        if value is not None: row[key] = value
                    row = {**row.pop("extra", {}), **row}
                    if columns is FIELD_COLUMNS:
                        row.setdefault("value_type", "float"); row.setdefault("unit", "1")
                        if "lower" in row or "upper" in row:
                            if not {"lower", "upper"} <= row.keys(): raise ValidationError("下限和上限需同时填写")
                            row["bounds"] = [row.pop("lower"), row.pop("upper")]
                        rule = row.get("execution_rounding", {})
                        for key in ("mode", "digits", "step", "origin"):
                            if f"execution_{key}" in row: rule[key] = row.pop(f"execution_{key}")
                        if rule: row["execution_rounding"] = rule
                        from experiment_planner.domain.template import parse_value
                        if "values" in row: row["values"] = [parse_value(v, row) for v in row["values"]]
                        if "fixed_value" in row: row["fixed_value"] = parse_value(row["fixed_value"], row)
                        if section == "parameters" and row["value_type"] == "text": row["use_as_model_input"] = False
                        if section != "parameters":
                            row.setdefault("is_response", row["value_type"] not in ("text", "category"))
                            row.setdefault("model_base", section == "measurements")
                        if section == "measurements": row.setdefault("measured_at", "after")
                    if section == "objectives": row.setdefault("direction", "maximize")
                    data[section].append(row)
                except (ValueError, TypeError, KeyError) as exc:
                    raise ValidationError(f"{title} 第 {row_number} 行：{exc}") from exc
        if "模板设置" in workbook:
            first = header_row(workbook["模板设置"]) + 1
            for row_number, (key, value, *rest) in enumerate(workbook["模板设置"].iter_rows(min_row=first, values_only=True), first):
                if key is None: continue
                try:
                    key = str(key).strip()
                    if key.split(".")[0] in {s for s, _ in TABLES.values()}: raise ValidationError("字段与目标请在对应工作表填写")
                    kind = ALIASES.get(rest[0], rest[0]) if rest else None
                    if kind == "str": value = str(value) if value is not None else ""
                    elif kind == "bool": value = _parse("enabled", value)
                    elif isinstance(value, str):
                        try: value = json.loads(value)
                        except json.JSONDecodeError: value = ALIASES.get(value, value)
                    value = _parse(key, value) if key in INT_KEYS else value
                    parts = key.split("."); target = data
                    for part in parts[:-1]: target = target.setdefault(part, {})
                    target[parts[-1]] = value
                except (ValueError, TypeError, AttributeError) as exc:
                    raise ValidationError(f"模板设置 第 {row_number} 行：{exc}") from exc
        return normalize_draft(data)
    finally: workbook.close()


def read_template_draft(path):
    if Path(path).suffix.lower() == ".xlsx": return read_template_excel(path)
    with open(path, encoding="utf-8") as source: return normalize_draft(json.load(source))
