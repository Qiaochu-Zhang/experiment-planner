"""Shared prediction tables and Excel export for selected experimental points."""
import json
import math
from pathlib import Path

from experiment_planner.domain.errors import ValidationError


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def _field(template, name):
    fields = {f["name"]: f for f in template.fields}
    if name in fields:
        return fields[name]
    policy = template.data.get("ratio_policy", {})
    if name in ("raw_selectivity", "raw_selectivity_abs", policy.get("stable_metric_name")):
        original = fields.get(policy.get("original_metric"), {})
        prefix = "稳定化：" if name == policy.get("stable_metric_name") else "原始比值："
        return {**original, "label": prefix + original.get("label", name)}
    return {}


def _statistics(item):
    q = item.get("quantiles")
    valid = (isinstance(q, (list, tuple)) and len(q) == 3
             and all(type(v) in (float, int) and math.isfinite(v) for v in q))
    mean = item.get("mean")
    mean = mean if type(mean) in (float, int) and math.isfinite(mean) else None
    notice = "；".join(dict.fromkeys(filter(None, (item.get("moment_notice"),
                         item.get("notice"), item.get("summary_notice")))))
    if not valid:
        q = [None, None, None]
        notice = "；".join(filter(None, [notice, "预测区间不可用"]))
    return {"预测均值": mean, "预测中位数50%": q[1], "下限2.5%": q[0], "上限97.5%": q[2],
            "预测状态": "区间不可用" if not valid else "仅有限样本" if item.get("nonfinite_sample_count", 0) else "可用",
            "区间类型": item.get("interval_kind"), "总样本数": item.get("sample_count"),
            "有限样本数": item.get("finite_sample_count"), "非有限样本数": item.get("nonfinite_sample_count"),
            "说明": notice}


def prediction_rows(prediction, template):
    variants = [("输入确定", prediction)]
    if "nominal_prediction" in prediction:
        variants = [("输入确定", prediction["nominal_prediction"]), ("含输入执行波动", prediction)]
    for interpretation, variant in variants:
        common = {"计算口径": interpretation, "预测可行概率": variant.get("feasibility_probability"),
                  "随机种子": variant.get("seed")}
        for name, item in variant.items():
            if not isinstance(item, dict) or "quantiles" not in item:
                continue
            field = _field(template, name)
            yield {**common, "结果类别": "响应/指标", "指标": name,
                   "指标名称": field.get("label", name), "单位": field.get("unit", "1"),
                   "目标变换": None, "目标方向": None, "目标定义": None,
                   **_statistics(item), "公式": item.get("expression")}
        for index, item in enumerate(variant.get("objective_predictions", []), 1):
            definition = item.get("definition", {})
            name = definition.get("metric", "")
            field = _field(template, name)
            stats = _statistics(item)
            stats["区间类型"] = ("latent_response_with_input_uncertainty" if interpretation == "含输入执行波动"
                                 else "latent_response")
            yield {**common, "结果类别": f"目标{index}", "指标": name,
                   "指标名称": field.get("label", name), "单位": field.get("unit", "1"),
                   "目标变换": definition.get("transform", "identity"),
                   "目标方向": definition.get("direction"), "目标定义": _json(definition),
                   **stats, "公式": item.get("expression")}


def write_sheet(workbook, title, rows, headers=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    headers = headers or list(dict.fromkeys(key for row in rows for key in row))
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(key) for key in headers])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str): cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245978")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[1].height = 60
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = min(44, max(18, len(header) * 1.6))
    return sheet


def point_prediction_tables(report, template):
    """One overview row per input point, with all interpretations in the detail table."""
    conditions, predictions = report.get("conditions", []), report.get("predictions", [])
    sources = report.get("source_experiments")
    if not conditions or len(conditions) != len(predictions):
        raise ValidationError("请先完成全部实验点的批量预测")
    if sources is not None and len(sources) != len(conditions):
        raise ValidationError("来源实验与预测结果数量不一致")
    overview, details = [], []
    for index, (condition, prediction) in enumerate(zip(conditions, predictions), 1):
        source = sources[index-1] if sources is not None else {}
        common = {"点序号": index, "来源实验编号（复制时）": source.get("experiment_id"),
                  "来源记录唯一标识": source.get("record_uid"), "复制时数据版本": source.get("copied_revision"),
                  **{f"条件：{f.get('label', f['name'])} [{f['name']}] ({f.get('unit', '1')})": condition[f["name"]]
                     for f in template.parameters}}
        row = {**common, "计算口径": "含输入执行波动" if "nominal_prediction" in prediction else "输入确定",
               "预测可行概率": prediction.get("feasibility_probability")}
        point_details = list(prediction_rows(prediction, template))
        if not point_details:
            raise ValidationError(f"预测点 {index} 没有可导出的响应结果")
        for detail in point_details:
            details.append({**common, **detail})
            if detail["计算口径"] == row["计算口径"]:
                prefix = f"{detail['结果类别']}：{detail['指标名称']} [{detail['指标']}] ({detail['单位']})"
                for key in ("预测均值", "预测中位数50%", "下限2.5%", "上限97.5%", "预测状态", "说明"):
                    row[f"{prefix} / {key}"] = detail[key]
        overview.append(row)
    return overview, details


def export_point_predictions(report, template, path):
    """Write the cached result snapshot, without accessing or changing the project."""
    from openpyxl import Workbook

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("批量实验点预测导出请选择 Excel (.xlsx)")
    overview, details = point_prediction_tables(report, template)
    definitions = [{"内部名称": f["name"], "显示名称": f.get("label", f["name"]), "类别": kind,
                    "单位": f.get("unit", "1"), "类型": f.get("value_type", "float"),
                    "允许下限": (f.get("bounds") or [None, None])[0],
                    "允许上限": (f.get("bounds") or [None, None])[1],
                    "离散值/固定值": _json(f.get("values", f.get("fixed_value"))),
                    "公式": f.get("expression"), "备注": f.get("note", "")}
                   for kind, group in (("输入条件", template.parameters), ("原始测量", template.measurements),
                                       ("派生指标", template.metrics)) for f in group]
    metadata = [
        ("项目", report.get("project_name")), ("预测时数据版本", report["revision"]),
        ("预测时模板版本", report["template_version"]), ("模型", report["model"]),
        ("随机种子", report.get("seed")), ("实验点数量", len(overview)),
        ("计算方式", "使用预测时项目数据拟合一次所选模型，然后同时预测全部条件；未生成趋势图。"),
        ("条件来源", "保留复制时的实验编号、唯一标识与完整条件；相同条件的不同实验各保留一行。编号可能被复用，请结合唯一标识识别。"),
        ("结果说明", "所选历史实验仍按原规则参与训练；这些结果属于模型后验预测，不是独立留出验证或实测值。"),
        ("区间含义", "下限2.5%和上限97.5%为中央95%潜在响应预测区间，未额外叠加未来测量噪声，不是确定的最大/最小值。"),
        ("预测值", "均值、中位数与区间保留计算精度。原始比值不声明均值存在，均值留空；不可用区间留空并保留原因。"),
        ("计算口径", "主表列出本次采用的预测；启用输入执行波动时，预测明细同时保留输入确定与含波动的结果。"),
        ("快照说明", "导出使用本次已完成的预测和模板快照，不重新拟合；后续修改记录或模型设置不会改写此结果，请核对数据版本。"),
        ("保存说明", "批量条件与预测仅在当前窗口保留，关闭或切换项目后清空；请及时导出 Excel。预测不修改实验记录。"),
        ("目标定义", _json(template.data.get("objectives", []))),
        ("输入不确定度配置", _json(report.get("input_uncertainty"))),
        ("模型数据集", _json(report.get("model_datasets", {}))),
        ("提示", "\n".join(report.get("notices", []))),
        ("原始警告", "\n".join(report.get("diagnostics", []))),
    ]
    for index, prediction in enumerate(report["predictions"], 1):
        if prediction.get("input_uncertainty"):
            metadata.append((f"点{index}输入波动配置与采样", _json(prediction["input_uncertainty"])))
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        write_sheet(workbook, "实验点预测", overview)
        write_sheet(workbook, "预测明细", details)
        write_sheet(workbook, "字段说明", definitions)
        write_sheet(workbook, "预测说明", [{"项目": key, "内容": value} for key, value in metadata])
        workbook.save(path)
    finally:
        workbook.close()
