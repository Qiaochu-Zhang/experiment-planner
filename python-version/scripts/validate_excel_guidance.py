"""Regenerate guided Excel examples and verify their complete import roundtrip."""
import json
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import load_workbook

from experiment_planner import __version__
from experiment_planner.application.service import PlannerService
from experiment_planner.application.records import normalize_observations
from experiment_planner.application.template_editing import blank_template
from experiment_planner.domain.template import Template
from experiment_planner.io.exchange import export_import_template, preview_import
from experiment_planner.io.template_excel import write_template_excel
from experiment_planner.storage.project import Project


def main():
    definitions = ROOT / 'examples/template_definitions'
    imports = ROOT / 'examples/import_templates'
    evidence = ROOT / f'docs/evidence/v{__version__}'
    evidence.mkdir(parents=True, exist_ok=True)
    template = Template.read(definitions / '模板定义示例.xlsx')
    template.write(definitions / '模板定义示例.xlsx')
    write_template_excel(blank_template(), definitions / '空白模板定义.xlsx')
    assert Template.read(definitions / '模板定义示例.xlsx').graph.evaluate({'quality': 3})['double_quality'].value == 6
    builtin = Template.builtin()
    empty = imports / '实验记录导入模板.xlsx'
    filled = imports / '实验记录填写示例_合成数据.xlsx'
    export_import_template(builtin, empty)
    export_import_template(builtin, filled)
    with tempfile.TemporaryDirectory() as folder:
        with Project.create(Path(folder) / 'examples.sqlite', '合成填写示例', builtin) as project:
            service = PlannerService(project)
            old = preview_import(service, imports / '实验记录填写示例_合成数据.csv')
            assert not old.errors and len(old.records) == 2
            old.commit(service)
            workbook = load_workbook(filled); sheet = workbook.active
            headers = [cell.value for cell in sheet[2]]
            for record, original in zip(project.experiments(), old.records, strict=True):
                row = {'external_id': original['external_id'], 'status': record['status'], 'note': record['note'], **record['actual']}
                for name, cell in record['observations'].items():
                    uncertainty = cell['uncertainty'] or {}
                    row.update({name: cell['value'], name + '_unit': cell['unit'], name + '_note': cell['note'],
                                name + '_uncertainty': uncertainty.get('amount'), name + '_uncertainty_kind': uncertainty.get('kind'),
                                name + '_distribution': uncertainty.get('distribution'), name + '_coverage_factor': uncertainty.get('coverage_factor'),
                                name + '_confidence': uncertainty.get('confidence')})
                sheet.append([row.get(key) for key in headers])
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str): cell.data_type = 's'
            workbook.save(filled); workbook.close()
            new = preview_import(service, filled)
            assert not new.errors and len(new.records) == len(old.records)
            for current, original in zip(new.records, old.records, strict=True):
                assert current['conditions'] == original['conditions']
                assert current['external_id'] == original['external_id']
                a = normalize_observations(builtin, current['observations'])
                b = normalize_observations(builtin, original['observations'])
                assert {n: (c['value'], c['uncertainty']) for n, c in a.items()} == {n: (c['value'], c['uncertainty']) for n, c in b.items()}
            assert not preview_import(service, empty).records
    artifacts = []
    for path in (empty, filled, definitions / '空白模板定义.xlsx', definitions / '模板定义示例.xlsx'):
        workbook = load_workbook(path)
        artifacts.append({'file': str(path.relative_to(ROOT)), 'sheets': workbook.sheetnames,
                          'validations': sum(len(sheet.data_validations.dataValidation) for sheet in workbook),
                          'hint_rows': [sheet.title for sheet in workbook if str(sheet['A1'].value).startswith('【填写提示】')]})
        workbook.close()
    result = {'version': __version__, 'python': sys.version, 'definition_roundtrip': 'passed',
              'legacy_csv_and_guided_excel_equal': 'passed', 'blank_record_template': 'passed', 'artifacts': artifacts}
    (evidence / 'excel-check.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__': main()
