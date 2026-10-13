import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog

from experiment_planner.storage.project import Project
from experiment_planner.storage.session import ProjectSession
from experiment_planner.ui.desktop import PythonWindow

pytestmark = pytest.mark.gui


def batch_for(window, conditions):
    project = window.require_project()
    return {"revision": project.revision, "template_version": project.template.data["template_version"],
            "request": {"n": 1}, "stage": "initialization", "notices": [],
            "candidates": [{"conditions": conditions, "arrangement": "initialization", "prediction": None}]}


def test_completed_recommendation_enables_real_button_and_exports(qt_application, tmp_path, conditions, monkeypatch):
    window = PythonWindow()
    assert not window.export_batch_button.isEnabled()
    window.set_project(Project.create(tmp_path / "ui.sqlite", "推荐导出"))
    assert not window.export_batch_button.isEnabled()
    result = batch_for(window, conditions)
    class CompletedTask:
        def poll(self): return {"result": result}
    window.task = CompletedTask()
    window.poll()
    assert window.export_batch_button.isEnabled()
    window.go_to(1); window.show(); qt_application.processEvents()
    assert window.export_batch_button.isVisible()
    destination = tmp_path / "推荐结果"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), "Excel (*.xlsx)"))
    window.export_batch_button.click()
    wb = load_workbook(destination.with_suffix(".xlsx"))
    assert wb.active.max_row == 2
    wb.close()
    assert "已导出批次 1" in window.statusBar().currentMessage()
    window.close_project()
    assert not window.export_batch_button.isEnabled()
    window.set_project(Project.create(tmp_path / "empty.sqlite", "空项目"))
    assert not window.export_batch_button.isEnabled()
    window.close()


def test_cancel_export_preserves_project_and_file(qt_application, tmp_path, conditions, monkeypatch):
    window = PythonWindow(); window.set_project(Project.create(tmp_path / "cancel.sqlite", "取消"))
    window.service.commit_batch(batch_for(window, conditions)); window.refresh()
    before = window.require_project().snapshot()
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    window.export_batch_button.click()
    assert window.require_project().snapshot() == before
    assert not list(tmp_path.glob("*.xlsx"))
    window.close()


def test_saved_session_can_export_after_reopen(qt_application, tmp_path, conditions, monkeypatch):
    path = tmp_path / "saved.sqlite"
    window = PythonWindow(); window.set_project(ProjectSession(path, creating=True))
    window.service.commit_batch(batch_for(window, conditions)); window.refresh()
    assert window.require_project().dirty
    window.save_project(); window.close_project()
    window.set_project(ProjectSession(path))
    assert window.export_batch_button.isEnabled()
    target = tmp_path / "重开.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), "Excel (*.xlsx)"))
    window.export_batch_button.click()
    assert target.is_file() and not window.require_project().dirty
    window.close()
