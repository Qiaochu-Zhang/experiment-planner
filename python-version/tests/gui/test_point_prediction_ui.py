import copy
import time

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QItemSelectionModel
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QFileDialog

from experiment_planner.domain.errors import ValidationError
from experiment_planner.storage.project import Project
from experiment_planner.ui.desktop import PythonWindow

pytestmark = pytest.mark.gui


@pytest.fixture
def window(qt_application, tmp_path, conditions, measurements, monkeypatch):
    window = PythonWindow()
    window.set_project(Project.create(tmp_path / "points.sqlite", "实验点预测"))
    window.service.add_record(conditions, measurements)
    window.service.add_record({**conditions, "etch_time_s": conditions["etch_time_s"]+.123456789}, measurements)
    window.service.add_record(conditions, {}, status="pending")
    window.refresh()
    failures = []
    monkeypatch.setattr(window, "show_failure", lambda *args: failures.append(args))
    window.test_failures = failures
    yield window
    window.close()


def select_rows(window, indices):
    selection = window.table.selectionModel()
    selection.clearSelection()
    for row in indices:
        selection.select(window.table.model().index(row, 0),
                         QItemSelectionModel.SelectionFlag.Select | QItemSelectionModel.SelectionFlag.Rows)


def report_for(window, request):
    return {"project_name": "实验点预测", "revision": window.require_project().revision,
            "template_version": window.require_project().template.data["template_version"],
            "model": request["model"], "seed": request["seed"],
            "conditions": copy.deepcopy(request["conditions"]),
            "source_experiments": copy.deepcopy(request["source_experiments"]),
            "predictions": [{"sio2_loss_nm": {"mean": i+.123456789123,
                "quantiles": [i-.5, i+.1, i+.5], "interval_kind": "latent_response"},
                "feasibility_probability": .75} for i, _ in enumerate(request["conditions"], 1)]}


@pytest.mark.parametrize("entry", ["records", "panel", "old_copy"])
def test_copy_selected_rows_in_table_order_and_keep_duplicate_conditions(window, entry):
    select_rows(window, [2, 0])
    before = window.require_project().snapshot()
    if entry == "records": window.copy_points_button.click()
    elif entry == "panel": window.point_prediction.copy_button.click()
    else: window.copy_conditions()
    panel = window.point_prediction
    assert not window.test_failures
    assert [p["source"]["experiment_id"] for p in panel.points] == [1, 3]
    assert panel.points[0]["conditions"] == panel.points[1]["conditions"]
    assert panel.table.rowCount() == 2
    assert window.prediction_tabs.currentWidget() is panel
    assert window.tabs.currentIndex() == 3 and window.navigation.currentIndex() == 1
    assert panel.predict_button.isEnabled() and not panel.export_button.isEnabled()
    assert window.require_project().snapshot() == before


def test_copy_without_selection_and_predict_without_points_are_actionable(window):
    window.table.clearSelection()
    with pytest.raises(ValidationError, match="Ctrl 或 Shift"): window.copy_point_conditions()
    with pytest.raises(ValidationError, match="复制"): window.start_point_prediction()
    with pytest.raises(ValidationError, match="先完成"): window.export_point_prediction_data()
    assert not window.point_prediction.predict_button.isEnabled()


def test_batch_request_uses_all_raw_conditions_current_model_seed_and_no_plot(window, monkeypatch):
    from experiment_planner.worker import process
    tasks = []
    class Task:
        def __init__(self, project, request, operation):
            self.request, self.operation = request, operation
            tasks.append(self)
        def cancel(self): self.cancelled = True
    monkeypatch.setattr(process, "CalculationTask", Task)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: pytest.fail("Batch prediction must not ask for plot destination"))
    select_rows(window, [1, 0])
    window.copy_point_conditions()
    window.model.setCurrentIndex(window.model.findData("gp_matern25_v1"))
    window.seed.setValue(17)
    window.plot_variables.setText("cl2_sccm,rf_w")
    panel = window.point_prediction
    assert "Matérn" in panel.model_hint.text() and "17" in panel.model_hint.text()
    before = window.require_project().snapshot()
    panel.predict_button.click()
    window.analysis_timer.stop()
    assert not window.test_failures
    [task] = tasks
    assert task.operation == "analyze" and task.request["model"] == "gp_matern25_v1"
    assert task.request["seed"] == 17
    assert task.request["conditions"] == [r["actual"] for r in window.require_project().experiments()[:2]]
    assert "plot_path" not in task.request and "variables" not in task.request
    assert not panel.copy_button.isEnabled() and not panel.predict_button.isEnabled()
    assert not window.analyze_button.isEnabled() and panel.cancel_button.isEnabled()
    with pytest.raises(ValidationError, match="等待或取消"): window.copy_point_conditions()
    assert window.require_project().snapshot() == before
    panel.cancel_button.click()
    assert task.cancelled and window.analysis_task is None
    assert panel.predict_button.isEnabled() and not panel.export_button.isEnabled()
    assert "已取消" in panel.output.toPlainText()


def test_completion_export_and_later_record_changes_keep_snapshot(window, monkeypatch, tmp_path):
    from experiment_planner.worker import process
    class Task:
        def __init__(self, project, request, operation): self.result = report_for(window, request)
        def poll(self): return {"result": self.result}
    monkeypatch.setattr(process, "CalculationTask", Task)
    select_rows(window, [0, 1]); window.copy_points_button.click()
    panel = window.point_prediction
    panel.predict_button.click(); window.poll_analysis()
    assert panel.export_button.isEnabled()
    assert panel.table.rowCount() == 2 and panel.report is not None
    assert all("来源记录唯一标识" not in panel.table.horizontalHeaderItem(i).text() for i in range(panel.table.columnCount()))
    assert any("预测均值" in panel.table.horizontalHeaderItem(i).text() for i in range(panel.table.columnCount()))
    original = copy.deepcopy(panel.report)
    original_uid = original["source_experiments"][0]["record_uid"]
    window.service.delete_records([1])
    old_record = window.require_project().experiment(2)
    assert window.service.add_record(old_record["actual"], {"sio2_initial_nm": 4.}) == 1
    window.refresh()
    assert "旧数据版本" in panel.output.toPlainText()
    destination = tmp_path / "预测结果"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), "Excel (*.xlsx)"))
    before = window.require_project().snapshot()
    panel.export_button.click()
    assert not window.test_failures
    assert window.require_project().snapshot() == before
    assert panel.report == original
    workbook = load_workbook(destination.with_suffix(".xlsx"))
    try:
        data = list(workbook["实验点预测"].values)
        row = dict(zip(data[0], data[1]))
        assert row["来源记录唯一标识"] == original_uid != window.require_project().experiment(1)["record_uid"]
        assert "已导出 2 个实验点" in window.statusBar().currentMessage()
    finally:
        workbook.close()


@pytest.mark.parametrize("result", [{"error": "模型拟合失败", "diagnostic": "测试原始诊断"},
                                   RuntimeError("poll failed")])
def test_failure_restores_controls_preserves_conditions_and_exposes_diagnostic(window, monkeypatch, result):
    from experiment_planner.worker import process
    class Task:
        def __init__(self, *a, **k): pass
        def poll(self):
            if isinstance(result, Exception): raise result
            return result
        def cancel(self): pass
    monkeypatch.setattr(process, "CalculationTask", Task)
    select_rows(window, [0, 1]); window.copy_point_conditions()
    panel = window.point_prediction
    before = window.require_project().snapshot()
    panel.predict_button.click(); window.poll_analysis()
    assert panel.table.rowCount() == 2 and panel.report is None
    assert panel.predict_button.isEnabled() and not panel.export_button.isEnabled()
    assert window.analysis_task is None
    assert window.require_project().snapshot() == before
    assert window.test_failures[0][2] == "批量实验点预测未完成"
    assert "原始诊断" in panel.output.toPlainText()


def test_export_cancel_remove_clear_replace_and_project_switch(window, monkeypatch, tmp_path):
    select_rows(window, [0, 1, 2]); window.copy_point_conditions()
    panel = window.point_prediction
    request = panel.request(window.require_project().template)
    request.update(model="gp_rbf_v1", seed=0)
    panel.display_report(report_for(window, request), window.require_project().template, window.require_project().revision)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    panel.export_button.click()
    assert panel.report is not None and not list(tmp_path.glob("*.xlsx"))
    panel.table.selectRow(1); panel.remove_button.click()
    assert [p["source"]["experiment_id"] for p in panel.points] == [1, 3]
    assert panel.report is None and not panel.export_button.isEnabled()
    select_rows(window, [1]); panel.copy_button.click()
    assert [p["source"]["experiment_id"] for p in panel.points] == [2]
    panel.clear_button.click()
    assert not panel.points and panel.table.rowCount() == 0
    select_rows(window, [0]); window.copy_point_conditions()
    window.close_project()
    assert not panel.points and panel.report is None
    window.set_project(Project.create(tmp_path / "other.sqlite", "其他项目"))
    assert not panel.points and not panel.predict_button.isEnabled()


def test_template_change_clears_copied_points(window):
    select_rows(window, [0]); window.copy_point_conditions()
    template = window.require_project().template
    parameters = template.parameters
    parameters[0]["label"] = "新显示名称"
    window.service.update_template(template.revised(parameters=parameters)); window.refresh()
    assert not window.point_prediction.points and window.point_prediction.report is None


@pytest.mark.model
def test_real_spawn_worker_batch_prediction_and_excel_via_buttons(window, qt_application, monkeypatch, tmp_path):
    # Real GP fit, spawn transport, timer polling, display and save dialog callback.
    window.model.setCurrentIndex(window.model.findData("gp_rbf_v1"))
    select_rows(window, [0, 1, 2]); window.copy_points_button.click()
    panel = window.point_prediction
    before = window.require_project().snapshot()
    panel.predict_button.click()
    deadline = time.monotonic() + 60
    while window.analysis_task is not None and time.monotonic() < deadline:
        qt_application.processEvents()
        window.poll_analysis()
        QTest.qWait(25)
    assert window.analysis_task is None, "Worker did not finish within 60 seconds"
    assert not window.test_failures
    assert panel.report is not None and len(panel.report["predictions"]) == 3
    assert "plot_path" not in panel.report
    destination = tmp_path / "真实GP.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), "Excel (*.xlsx)"))
    panel.export_button.click()
    assert destination.exists() and window.require_project().snapshot() == before
    workbook = load_workbook(destination)
    try:
        assert workbook.active.max_row == 4
    finally:
        workbook.close()
