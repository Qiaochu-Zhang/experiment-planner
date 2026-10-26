import copy

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog, QDialog, QDialogButtonBox, QPlainTextEdit, QMessageBox

from experiment_planner.application.importance import analyze_importance
from experiment_planner.domain.errors import ValidationError
from experiment_planner.io.exchange import export_records
from experiment_planner.storage.project import Project
from experiment_planner.ui.data_quality import ExperimentBatchDialog, DataWeightingDialog
from experiment_planner.ui.desktop import PythonWindow
from experiment_planner.ui.messages import record_result_text
from experiment_planner.ui.window import RecordDialog

pytestmark = pytest.mark.gui


@pytest.fixture
def quality_window(qt_application, tmp_path, conditions, measurements):
    window = PythonWindow(); window.set_project(Project.create(tmp_path / "quality-ui.sqlite", "实验质量界面"))
    window.service.add_record(conditions, measurements, performed_at="2026-01-01")
    window.service.add_record({**conditions, "etch_time_s": 160.}, {**measurements, "sio2_remaining_nm": 70.}, performed_at="2026-02-01")
    window.refresh()
    yield window
    window.close()


def test_batch_assignment_ui_and_weight_dialog_roundtrip(quality_window, monkeypatch):
    window = quality_window; p = window.require_project()
    window.table.selectRow(0)
    def execute(dialog):
        dialog.batch.setCurrentText("机台调整前"); dialog.confidence.setValue(.25); dialog.date.setText("2025-12-31")
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ExperimentBatchDialog, "exec", execute)
    window.batch_button.click()
    assert p.experiment(1)["experiment_batch"] == "机台调整前"
    assert p.experiment(1)["performed_at"].startswith("2025-12-31")
    assert p.template.data["data_weighting"]["batch_confidence"] == {"机台调整前": .25}
    def weighting(dialog):
        assert dialog.editors["机台调整前"].value() == .25
        dialog.half_life.setValue(30); dialog.editors["机台调整前"].setValue(.5)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(DataWeightingDialog, "exec", weighting)
    window.weighting_button.click()
    assert p.template.data["data_weighting"] == {"recency_half_life_days": 30., "batch_confidence": {"机台调整前": .5}}
    headers = [window.table.horizontalHeaderItem(i).text() for i in range(window.table.columnCount())]
    assert window.table.item(0, headers.index("批次可信度")).text() == "0.5"
    window.service.undo(); window.service.undo(); window.refresh()
    assert p.experiment(1)["experiment_batch"] == ""


def test_batch_and_weight_cancel_preserve_project(quality_window, monkeypatch):
    window = quality_window; before = window.require_project().snapshot()
    window.table.selectRow(0)
    monkeypatch.setattr(ExperimentBatchDialog, "exec", lambda self: QDialog.DialogCode.Rejected)
    monkeypatch.setattr(DataWeightingDialog, "exec", lambda self: QDialog.DialogCode.Rejected)
    window.batch_button.click(); window.weighting_button.click()
    assert window.require_project().snapshot() == before
    window.table.clearSelection()
    with pytest.raises(ValidationError, match="Ctrl"): window.assign_experiment_batch()


def test_pending_record_dialog_preserves_status_and_date(quality_window, conditions):
    window = quality_window; p = window.require_project()
    eid = window.service.add_record(conditions, {}, status="pending", experiment_batch="待做批")
    dialog = RecordDialog(p.template, p.experiment(eid))
    assert dialog.values()["status"] == "pending"
    assert dialog.values()["experiment_batch"] == "待做批"
    assert not dialog.values()["performed_at"]
    window.service.revise_record(eid, **dialog.values())
    assert p.experiment(eid)["performed_at"] is None
    dialog.close()


def test_records_show_prediction_preserve_actual_and_hide_changed_conditions(quality_window, conditions, measurements):
    window = quality_window; p = window.require_project()
    batch = {"revision": p.revision, "template_version": p.template.data["template_version"], "request": {"n": 1},
        "stage": "model_driven", "notices": [], "candidates": [{"conditions": conditions, "arrangement": "new", "prediction": {
            "sio2_loss_nm": {"mean": 12.3456, "quantiles": [1., 12., 20.]},
            "selectivity": {"mean": None, "quantiles": [1., 4.5, 9.]}}}]}
    _, [eid] = window.service.commit_batch(batch); window.refresh()
    headers = [window.table.horizontalHeaderItem(i).text() for i in range(window.table.columnCount())]
    column = next(i for i, h in enumerate(headers) if h.startswith("SiO2刻蚀量"))
    assert window.table.item(2, column).text() == "未测（预测均值 12.346）"
    assert "预测中位数 4.5" in record_result_text(p.experiment(eid), p.template.metrics[-1], p.template)
    assert not p.experiment(eid)["observations"]
    window.service.revise_record(eid, {**conditions, "etch_time_s": 50.}, {}, status="pending")
    window.refresh(); assert "原预测不适用" in window.table.item(2, column).text()
    window.service.revise_record(eid, conditions, measurements, status="completed")
    window.refresh(); assert window.table.item(2, column).text() == "10.0"
    assert p.experiment(eid)["experiment_batch"].startswith("推荐批次")


@pytest.mark.parametrize("mean,quantiles,expected", [(0., [0., 1., 2.], "预测均值 0"), (None, [0., 0., 0.], "预测中位数 0"),
    (float("nan"), None, "预测不可用"), (None, None, "预测不可用")])
def test_pending_prediction_zero_and_unavailable_are_explicit(quality_window, mean, quantiles, expected):
    p = quality_window.require_project(); field = p.template.metrics[0]
    record = {"status": "pending", "actual": {}, "observations": {}, "derived": {},
              "prediction": {field["name"]: {"mean": mean, "quantiles": quantiles}}}
    assert expected in record_result_text(record, field, p.template)


def test_import_ui_defaults_to_update_and_previews_then_commits(quality_window, tmp_path, monkeypatch):
    window = quality_window; p = window.require_project()
    path = tmp_path / "bulk.xlsx"; export_records(p, path)
    wb = load_workbook(path); columns = {c.value: c.column for c in wb.active[1]}
    wb.active.cell(2, columns["sio2_remaining_nm"], 50.); wb.save(path); wb.close()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(path), ""))
    messages = []; monkeypatch.setattr(QMessageBox, "information", lambda *a: messages.append(a[-1]))
    def execute(dialog):
        output = dialog.findChild(QPlainTextEdit).toPlainText()
        assert "将修改 1 条" in output and "SiO2剩余厚度" in output and "90.0 → 50.0" in output
        dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Save).click()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(QDialog, "exec", execute)
    window.import_data()
    assert len(p.experiments()) == 2 and p.experiment(1)["derived"]["sio2_loss_nm"]["value"] == 50.
    assert "修改 1 条" in messages[0]
    window.service.undo(); assert p.experiment(1)["observations"]["sio2_remaining_nm"]["value"] == 90.


def test_importance_ui_worker_request_display_stale_export_and_close(quality_window, tmp_path, monkeypatch):
    window = quality_window; p = window.require_project()
    window.model.setCurrentIndex(window.model.findData("bayesian_linear_v1")); window.seed.setValue(13)
    report = analyze_importance(p.snapshot(), {"model": "bayesian_linear_v1", "seed": 13})
    tasks = []
    class Task:
        def __init__(self, project, request, operation):
            assert request == {"model": "bayesian_linear_v1", "seed": 13} and operation == "importance"
            tasks.append(self)
        def poll(self): return {"result": report}
        def cancel(self): self.cancelled = True
    from experiment_planner.worker import process
    monkeypatch.setattr(process, "CalculationTask", Task)
    window.importance.run_button.click(); window.analysis_timer.stop()
    assert len(tasks) == 1 and not window.analyze_button.isEnabled() and window.importance.cancel_button.isEnabled()
    window.poll_analysis()
    assert window.importance.table.rowCount() == len(p.template.parameters)
    assert window.importance.export_button.isEnabled() and window.analyze_button.isEnabled()
    before = p.snapshot()
    target = tmp_path / "importance.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    window.importance.export_button.click(); assert target.is_file()
    assert p.snapshot() == before
    window.service.assign_experiment_batch([1], "旧", confidence=.3); window.refresh()
    assert window.importance.output.toPlainText().startswith("项目已更改")
    assert window.model.currentData() == "bayesian_linear_v1" and window.seed.value() == 13
    window.importance.run_button.click(); window.analysis_timer.stop(); window.cancel_analysis()
    assert tasks[-1].cancelled and "已取消" in window.importance.output.toPlainText()
    window.close_project(); assert window.importance.report is None


def test_importance_failure_uses_existing_diagnostics(quality_window, monkeypatch):
    window = quality_window
    window.analysis_kind = "importance"
    class Task:
        def poll(self): return {"error": "合成失败", "diagnostic": "原始诊断"}
    window.analysis_task = Task(); window.importance.set_busy(True)
    failures = []; monkeypatch.setattr(window, "show_failure", lambda *args: failures.append(args))
    window.poll_analysis()
    assert failures[0][1] == "原始诊断"
    assert "合成失败" in window.importance.output.toPlainText() and window.importance.run_button.isEnabled()


def test_stale_weighting_dialog_cannot_overwrite_newer_settings(quality_window, monkeypatch):
    window = quality_window; p = window.require_project()
    def execute(dialog):
        window.service.assign_experiment_batch([1], "新增批次", confidence=.2)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(DataWeightingDialog, "exec", execute)
    with pytest.raises(ValidationError, match="项目已修改"): window.edit_data_weighting()
    assert p.template.data["data_weighting"]["batch_confidence"] == {"新增批次": .2}
