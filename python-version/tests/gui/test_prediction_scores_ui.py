from copy import deepcopy
import time

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QItemSelectionModel
from PySide6.QtWidgets import QDialog, QFileDialog, QPlainTextEdit

from experiment_planner.domain.errors import ValidationError
from experiment_planner.io.score_export import prediction_score_report
from experiment_planner.storage.project import Project
from experiment_planner.ui.desktop import PythonWindow
from experiment_planner.ui.prediction_scores import PredictionScoresDialog
from experiment_planner.ui.window import RecordDialog

pytestmark = pytest.mark.gui


@pytest.fixture
def score_window(qt_application, tmp_path, conditions, measurements, monkeypatch):
    window=PythonWindow(); window.set_project(Project.create(tmp_path/"score-ui.sqlite", "预测评分界面"))
    p=window.require_project()
    window.service.update_template(p.template.revised(ratio_policy={**p.template.data["ratio_policy"], "mode":"remove"}))
    window.service.add_record({**conditions,"etch_time_s":20.},measurements)
    window.service.add_record({**conditions,"etch_time_s":160.},{**measurements,"sio2_remaining_nm":60.})
    window.service.add_record(conditions,{},status="pending"); window.refresh()
    window.score_failures=[]
    monkeypatch.setattr(window,"show_failure",lambda *args:window.score_failures.append(args))
    yield window
    window.close()


def select(window, indices):
    selection=window.table.selectionModel(); selection.clearSelection()
    for row in indices:
        selection.select(window.table.model().index(row,0),QItemSelectionModel.SelectionFlag.Select|QItemSelectionModel.SelectionFlag.Rows)


def prediction_report(project, request):
    return {"revision":project.revision,"template_version":project.template.data["template_version"],
            "model":request["model"],"seed":request["seed"],"conditions":deepcopy(request["conditions"]),
            "source_experiments":deepcopy(request["source_experiments"]),
            "predictions":[{"sio2_loss_nm":{"mean":10.,"quantiles":[9.,10.,11.]}} for _ in request["conditions"]]}


def fake_worker(window, monkeypatch):
    tasks=[]
    class Task:
        def __init__(self, project, request, operation):
            assert operation=="analyze" and "plot_path" not in request
            self.request=request; self.result={"result":prediction_report(project,request)};tasks.append(self)
        def poll(self): return self.result
        def cancel(self): self.cancelled=True
    from experiment_planner.worker import process
    monkeypatch.setattr(process,"CalculationTask",Task)
    return tasks


def test_record_prediction_button_saves_exact_selected_sources_and_backfill_scores(score_window, monkeypatch, measurements):
    window=score_window; p=window.require_project(); tasks=fake_worker(window,monkeypatch)
    select(window,[2]); window.model.setCurrentIndex(window.model.findData("bayesian_linear_v1")); window.seed.setValue(17)
    before=p.experiment(3)
    window.record_prediction_button.click(); window.analysis_timer.stop()
    assert not window.record_prediction_button.isEnabled() and window.analysis_kind=="record_predictions"
    assert tasks[0].request["conditions"]==[before["actual"]] and tasks[0].request["seed"]==17
    window.poll_analysis()
    assert not window.score_failures and window.record_prediction_button.isEnabled()
    snapshot=deepcopy(p.experiment(3)["last_prediction"])
    assert snapshot["prediction"]["sio2_loss_nm"]["mean"]==10.
    dialog=RecordDialog(p.template,p.experiment(3))
    for name,value in measurements.items(): dialog.cells[name][0].setText(str(value))
    dialog.cells["sio2_remaining_nm"][0].setText("87")
    dialog.status.setCurrentIndex(dialog.status.findData("completed"))
    window.service.revise_record(3,**dialog.values()); window.refresh(); dialog.close()
    assert p.experiment(3)["last_prediction"]==snapshot
    [score]=p.experiment(3)["prediction_evaluation"]["rows"]
    assert score["absolute_error"]==3. and score["interval_score95"]==82.
    headers=[window.table.horizontalHeaderItem(i).text() for i in range(window.table.columnCount())]
    summary=window.table.item(2,headers.index("上次预测 / 评分（越小越好）")).text()
    assert "AE=3" in summary and "IS95=82" in summary


def test_record_all_without_selection_and_cancel_retains_previous_data(score_window,monkeypatch):
    window=score_window; p=window.require_project(); tasks=fake_worker(window,monkeypatch)
    before=p.snapshot(); window.table.clearSelection(); window.record_prediction_button.click();window.analysis_timer.stop()
    assert len(tasks[0].request["source_experiments"])==3
    window.cancel_analysis()
    assert tasks[0].cancelled and p.snapshot()==before and "已取消" in window.record_prediction_status.text()


@pytest.mark.parametrize("failure",["worker","stale","identity"])
def test_failed_recording_retains_prior_prediction_and_exposes_diagnostic(score_window,monkeypatch,failure):
    window=score_window;p=window.require_project();tasks=fake_worker(window,monkeypatch)
    select(window,[2]);window.record_prediction_button.click();window.analysis_timer.stop()
    if failure=="worker": tasks[0].result={"error":"合成计算失败","diagnostic":"原始错误"}
    elif failure=="stale": window.service.assign_experiment_batch([3],"新批次")
    else: tasks[0].result["result"]["source_experiments"][0]["record_uid"]="wrong"
    before=p.snapshot();window.poll_analysis()
    assert window.score_failures and p.snapshot()==before and "last_prediction" not in p.experiment(3)
    assert window.record_prediction_button.isEnabled() and "原预测保留" in window.record_prediction_status.text()


def test_normal_batch_prediction_stays_readonly_until_record_button_clicked(score_window,monkeypatch):
    window=score_window;p=window.require_project();fake_worker(window,monkeypatch)
    select(window,[0,2]);window.copy_point_conditions();before=p.snapshot()
    window.point_prediction.predict_button.click();window.analysis_timer.stop();window.poll_analysis()
    assert p.snapshot()==before and window.point_prediction.record_button.isEnabled()
    window.point_prediction.record_button.click()
    assert not window.score_failures and not window.point_prediction.record_button.isEnabled()
    assert p.experiment(1)["last_prediction"] and p.experiment(3)["last_prediction"]
    assert p.experiment(1)["prediction_evaluation"]["rows"][0]["comparison_kind"]=="预测时已有结果对照"


def test_score_dialog_formulas_export_and_selected_entry_filter(score_window,monkeypatch,tmp_path):
    window=score_window;p=window.require_project();fake_worker(window,monkeypatch)
    select(window,[2]);window.record_prediction_button.click();window.analysis_timer.stop();window.poll_analysis()
    row=p.experiment(3);window.service.revise_record(3,row["actual"],{"sio2_initial_nm":100.,"sio2_remaining_nm":87.},status="partial")
    window.refresh();select(window,[2])
    seen=[]
    def execute(dialog):
        assert [e["record"]["id"] for e in dialog.report["entries"]]==[3]
        seen.append(dialog.report);return QDialog.DialogCode.Rejected
    monkeypatch.setattr(PredictionScoresDialog,"exec",execute)
    window.prediction_scores_button.click();assert len(seen)==1
    report=prediction_score_report(p,[3]);dialog=PredictionScoresDialog(report)
    assert dialog.table.rowCount()==1 and "MAE=3" in dialog.summary.toPlainText()
    assert "82" in dialog.table.item(0,8).text()
    path=tmp_path/"score-ui.xlsx"
    monkeypatch.setattr(QFileDialog,"getSaveFileName",lambda *a,**k:(str(path),""))
    dialog.export_button.click();assert path.is_file()
    wb=load_workbook(path);assert "评分公式" in wb.sheetnames;wb.close()
    formulas=[]
    def formula_exec(child):
        formulas.append(child.findChild(QPlainTextEdit).toPlainText());return QDialog.DialogCode.Rejected
    monkeypatch.setattr(QDialog,"exec",formula_exec)
    dialog.formulas_button.click();assert "40×max" in formulas[0] and "测量" in formulas[0]
    dialog.close()


def test_real_worker_records_prediction_and_preserves_it_after_measurement(score_window,qt_application):
    window=score_window;p=window.require_project();select(window,[2])
    window.model.setCurrentIndex(window.model.findData("bayesian_linear_v1"))
    window.record_prediction_button.click()
    deadline=time.monotonic()+55
    while window.analysis_task is not None and time.monotonic()<deadline:
        qt_application.processEvents();time.sleep(.02)
    assert window.analysis_task is None and not window.score_failures
    snapshot=deepcopy(p.experiment(3)["last_prediction"])
    window.service.revise_record(3,p.experiment(3)["actual"],{"sio2_initial_nm":100.,"sio2_remaining_nm":88.},status="partial")
    assert p.experiment(3)["last_prediction"]==snapshot
    [score]=p.experiment(3)["prediction_evaluation"]["rows"][:1]
    assert score["measured"]==12. and score["absolute_error"] is not None


def test_empty_project_has_actionable_recording_error(qt_application,tmp_path):
    window=PythonWindow();window.set_project(Project.create(tmp_path/"empty-score.sqlite","空项目"))
    with pytest.raises(ValidationError,match="至少一条"):window.start_record_prediction()
    with pytest.raises(ValidationError,match="请先"):window.show_prediction_scores()
    window.close()
