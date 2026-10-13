import copy
import json

import pytest
from openpyxl import load_workbook

from experiment_planner.application.service import PlannerService
from experiment_planner.domain.errors import StaleVersionError, ValidationError
from experiment_planner.domain.template import Template
from experiment_planner.io.batch_export import export_batch
from experiment_planner.storage.project import Project, encode


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def batch_for(project, conditions, prediction=None):
    return {"revision": project.revision, "template_version": project.template.data["template_version"],
            "request": {"n": 1, "model": "bayesian_linear_v1", "seed": 3},
            "stage": "model_driven" if prediction else "initialization", "notices": [],
            "candidates": [{"conditions": conditions, "prediction": prediction, "arrangement": "new"}]}


def test_full_conditions_numeric_precision_and_literal_text(tmp_path):
    template = Template({"schema_version": 2, "template_version": 1,
        "parameters": [{"name": "x", "label": "流量", "unit": "sccm", "bounds": [0, 10], "display_digits": 0},
                       {"name": "fixed", "fixed_value": 5.0, "bounds": [0, 10]},
                       {"name": "gas", "value_type": "category", "values": ["A", "B"]},
                       {"name": "enabled", "value_type": "bool"},
                       {"name": "note", "value_type": "text", "use_as_model_input": False}],
        "measurements": [{"name": "y", "label": "=1+1", "unit": "nm", "is_response": True, "display_digits": 0}],
        "derived_metrics": [], "objectives": [{"metric": "y", "direction": "maximize"}],
        "optimization": {"mode": "single"}, "ratio_policy": {"mode": "remove"}})
    conditions = {"x": 1.23456789123, "fixed": 5., "gas": "B", "enabled": False, "note": "=1+1"}
    value = 9.87654321234
    stats = {"mean": value, "quantiles": [1.123456789, value, 15.7654321], "interval_kind": "latent_response"}
    prediction = {"y": stats, "feasibility_probability": .75, "seed": 3,
                  "objective_predictions": [{**stats, "definition": template.data["objectives"][0]}]}
    with Project.create(tmp_path / "mixed.sqlite", "=合成项目", template) as p:
        service = PlannerService(p)
        service.add_record(conditions, {"y": 4})  # Existing history is not exported.
        batch = batch_for(p, conditions, prediction)
        before_batch = copy.deepcopy(batch)
        bid, ids = service.commit_batch(batch)
        assert batch == before_batch
        before = p.snapshot(), p.history(), p.batches()
        export_batch(p, tmp_path / "完整.xlsx")
        assert (p.snapshot(), p.history(), p.batches()) == before
        with_workbook = load_workbook(tmp_path / "完整.xlsx")
        try:
            assert with_workbook.sheetnames == ["推荐实验", "预测明细", "字段说明", "批次说明"]
            [row] = rows(with_workbook["推荐实验"])
            assert row["实验编号（推荐时）"] == ids[0] == 2
            assert row["批次编号"] == bid
            for name, expected in conditions.items():
                header = next(k for k in row if k.startswith("条件：") and f"[{name}]" in k)
                assert row[header] == expected
            detail = rows(with_workbook["预测明细"])
            assert len(detail) == 2
            assert detail[0]["预测中位数50%"] == value
            assert detail[0]["下限2.5%"] == stats["quantiles"][0]
            assert detail[0]["上限97.5%"] == stats["quantiles"][2]
            assert detail[0]["预测可行概率"] == .75
            assert detail[1]["结果类别"] == "目标1"
            assert detail[1]["目标方向"] == "maximize"
            assert all(cell.data_type != "f" for sheet in with_workbook for line in sheet for cell in line)
            assert all(sheet.freeze_panes == "A2" for sheet in with_workbook)
        finally:
            with_workbook.close()


def test_reopen_export_keeps_original_ids_conditions_and_template(service, conditions, measurements, tmp_path):
    p = service.project
    batch = batch_for(p, conditions)
    bid, [eid] = service.commit_batch(batch)
    original = p.experiment(eid)
    new_conditions = {**conditions, "etch_time_s": conditions["etch_time_s"] + 1}
    service.revise_record(eid, new_conditions, measurements)
    service.delete_records([eid])
    assert service.add_record(new_conditions, measurements) == eid
    parameters = p.template.parameters
    parameters[0]["label"] = "新名称"
    service.update_template(p.template.revised(parameters=parameters))
    with Project(p.path) as reopened:
        export_batch(reopened, tmp_path / "历史.xlsx", batch_id=bid)
    wb = load_workbook(tmp_path / "历史.xlsx")
    try:
        [row] = rows(wb["推荐实验"])
        assert row["记录唯一标识"] == original["record_uid"] != p.experiment(eid)["record_uid"]
        time_key = next(k for k in row if "[etch_time_s]" in k)
        assert row[time_key] == conditions["etch_time_s"]
        assert "新名称" not in " ".join(row)
        assert dict((r["项目"], r["内容"]) for r in rows(wb["批次说明"]))["推荐时模板版本"] == batch["template_version"]
    finally:
        wb.close()


def test_legacy_batch_uses_audit_ids_and_explains_template_fallback(service, conditions, tmp_path):
    p = service.project
    original = batch_for(p, conditions)
    bid, ids = service.commit_batch(original)
    with p.transaction():
        p.db.execute("UPDATE batches SET payload=? WHERE id=?", (encode(original), bid))
    export_batch(p, tmp_path / "旧批次.xlsx")
    wb = load_workbook(tmp_path / "旧批次.xlsx")
    try:
        [row] = rows(wb["推荐实验"])
        assert row["实验编号（推荐时）"] == ids[0]
        assert row["记录唯一标识"] is None
        assert "暂无预测" in row["预测状态"]
        assert rows(wb["预测明细"])[0]["下限2.5%"] is None
        assert "旧批次未保存模板快照" in str(list(wb["批次说明"].values))
    finally:
        wb.close()


def test_export_latest_or_selected_batch_and_repeats(service, conditions, measurements, tmp_path):
    p = service.project
    source = service.add_record(conditions, measurements)
    first, _ = service.commit_batch(batch_for(p, conditions))
    second = batch_for(p, {**conditions, "etch_time_s": conditions["etch_time_s"] + 1})
    second["request"]["n"] = 2
    second["candidates"].insert(0, {"conditions": conditions, "arrangement": "repeat", "repeat_of": source})
    bid, ids = service.commit_batch(second)
    for selected, expected in ((None, ids), (first, [2])):
        path = tmp_path / f"batch-{selected}.xlsx"
        export_batch(p, path, batch_id=selected)
        wb = load_workbook(path)
        try:
            data = rows(wb.active)
            assert [r["实验编号（推荐时）"] for r in data] == expected
            if selected is None:
                assert data[0]["安排"] == "复测" and data[0]["复测来源编号"] == source
                assert data[1]["批次编号"] == bid
        finally:
            wb.close()


@pytest.mark.parametrize("unavailable", [False, True])
def test_nonfinite_and_raw_ratio_keep_blank_means_and_notices(service, conditions, tmp_path, unavailable):
    stats = {"mean": None, "quantiles": None if unavailable else [1., 2., 3.],
             "sample_count": 4, "finite_sample_count": 0 if unavailable else 3,
             "nonfinite_sample_count": 4 if unavailable else 1, "moment_notice": "原始比值不声明均值存在",
             "summary_notice": "全部不可用" if unavailable else "区间仅描述有限样本"}
    service.commit_batch(batch_for(service.project, conditions, {"raw_selectivity": stats}))
    export_batch(service.project, tmp_path / "比值.xlsx")
    wb = load_workbook(tmp_path / "比值.xlsx")
    try:
        [detail] = rows(wb["预测明细"])
        assert detail["预测均值"] is None
        assert detail["预测状态"] == ("区间不可用" if unavailable else "仅有限样本")
        assert detail["下限2.5%"] == (None if unavailable else 1.)
        assert detail["非有限样本数"] == (4 if unavailable else 1)
        assert "不声明均值存在" in detail["说明"]
    finally:
        wb.close()


@pytest.mark.parametrize("reason", ["no_batch", "empty", "missing_id", "extension"])
def test_invalid_export_does_not_create_file(service, conditions, tmp_path, reason):
    if reason != "no_batch":
        batch = batch_for(service.project, conditions)
        if reason == "empty": batch["candidates"] = []
        service.commit_batch(batch)
    path = tmp_path / ("output.csv" if reason == "extension" else "output.xlsx")
    with pytest.raises(ValidationError):
        export_batch(service.project, path, batch_id=999 if reason == "missing_id" else None)
    assert not path.exists()


def test_batch_snapshot_failure_rolls_back_and_stale_batch_is_rejected(service, conditions, monkeypatch):
    p = service.project
    batch = batch_for(p, conditions)
    before = p.snapshot(), p.batches(), p.history()
    with monkeypatch.context() as context:
        def fail(*args): raise RuntimeError("模拟审计保存失败")
        context.setattr(p, "log", fail)
        with pytest.raises(RuntimeError): service.commit_batch(batch)
    assert (p.snapshot(), p.batches(), p.history()) == before
    service.add_record(conditions, {})
    with pytest.raises(StaleVersionError): service.commit_batch(batch)
    assert p.batches() == []


@pytest.mark.model
@pytest.mark.parametrize("preset", ["bayesian_linear_v1", "gp_rbf_v1", "gp_matern25_v1"])
@pytest.mark.parametrize("with_uncertainty", [False, True])
def test_real_recommendation_export_matches_saved_posterior(tmp_path, preset, with_uncertainty):
    from experiment_planner.engine.planner import generate
    from experiment_planner.engine.space import BatchRequest
    template = Template({"schema_version": 2, "template_version": 1,
        "parameters": [{"name": "x", "bounds": [0., 1.],
                        "input_uncertainty": {"enabled": True, "kind": "bounds", "amount": .1}}],
        "measurements": [{"name": "y", "unit": "nm", "is_response": True}], "derived_metrics": [],
        "objectives": [{"metric": "y", "direction": "maximize"}],
        "optimization": {"mode": "single"}, "ratio_policy": {"mode": "remove"},
        "input_uncertainty": {"enabled": with_uncertainty, "samples": 8}})
    with Project.create(tmp_path / "real.sqlite", "合成推荐导出", template) as p:
        service = PlannerService(p)
        for x in (.1, .4, .8):
            service.add_record({"x": x}, {"y": {"value": 1+2*x, "uncertainty": {"kind": "std", "amount": .05}}})
        batch = generate(p.snapshot(), BatchRequest(n=2, repeat_ids=[1], model=preset, pool_size=4, seed=19))
        service.commit_batch(batch)
        export_batch(p, tmp_path / "真实模型.xlsx")
    wb = load_workbook(tmp_path / "真实模型.xlsx")
    try:
        data = rows(wb["预测明细"])
        assert len(rows(wb["推荐实验"])) == 2
        assert len(data) == (8 if with_uncertainty else 4)
        for index, candidate in enumerate(batch["candidates"], 1):
            expected = candidate["prediction"]
            for detail in [r for r in data if r["批内序号"] == index]:
                source = expected["nominal_prediction"] if with_uncertainty and detail["计算口径"] == "输入确定" else expected
                item = source["y"] if detail["结果类别"] == "响应/指标" else source["objective_predictions"][0]
                assert detail["预测均值"] == pytest.approx(item["mean"], rel=1e-14)
                assert [detail[k] for k in ("下限2.5%", "预测中位数50%", "上限97.5%")]==pytest.approx(item["quantiles"], rel=1e-14)
        if with_uncertainty:
            assert "输入波动配置与采样" in str(list(wb["批次说明"].values))
    finally:
        wb.close()
