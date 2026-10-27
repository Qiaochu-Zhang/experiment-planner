from copy import deepcopy
import csv
import json

import pytest
from openpyxl import load_workbook

from experiment_planner.application.analysis import analyze
from experiment_planner.application.service import PlannerService
from experiment_planner.domain.errors import ValidationError, StaleVersionError
from experiment_planner.domain.template import Template
from experiment_planner.io.exchange import export_records, preview_import
from experiment_planner.io.score_export import prediction_score_report, export_prediction_scores
from experiment_planner.storage.project import Project, encode
from experiment_planner.storage.session import ProjectSession


@pytest.fixture
def history_service(tmp_path):
    template = Template({"schema_version": 2, "template_version": 1, "parameters": [{"name": "x", "bounds": [0, 1]}],
        "measurements": [{"name": "y", "is_response": True}],
        "derived_metrics": [{"name": "double_y", "expression": "y * 2"}],
        "objectives": [{"metric": "y", "direction": "maximize"}], "optimization": {"mode": "single"},
        "ratio_policy": {"mode": "remove"}})
    project = Project.create(tmp_path / "history.sqlite", "预测回填合成测试", template)
    yield PlannerService(project)
    project.close()


def report_for(service, ids, mean=10.):
    p = service.project
    records = [p.experiment(eid) for eid in ids]
    return {"revision": p.revision, "template_version": p.template.data["template_version"],
            "model": "bayesian_linear_v1", "seed": 7, "conditions": [r["actual"] for r in records],
            "source_experiments": [{"experiment_id": r["id"], "record_uid": r.get("record_uid")} for r in records],
            "predictions": [{"y": {"mean": mean, "quantiles": [mean-1,mean,mean+1]},
                             "double_y": {"mean": 2*mean, "quantiles": [2*(mean-1),2*mean,2*(mean+1)]}} for r in records]}


def add_pending(service, x=.5):
    return service.add_record({"x": x}, {}, status="pending")


def test_manual_prediction_snapshot_backfill_corrections_and_undo_reopen(history_service, monkeypatch):
    service = history_service; p = service.project; eid = add_pending(service)
    assert service.save_record_predictions(report_for(service, [eid])) == [eid]
    snapshot = deepcopy(p.experiment(eid)["last_prediction"])
    from experiment_planner.engine import models
    monkeypatch.setattr(models, "fit_models", lambda *a, **k: pytest.fail("Backfill scoring must never fit"))
    service.revise_record(eid, {"x": .5}, {"y": 13.}, status="completed")
    record = p.experiment(eid)
    scores = {r["metric"]: r for r in record["prediction_evaluation"]["rows"]}
    assert scores["y"]["absolute_error"] == 3. and scores["y"]["interval_score95"] == 82.
    assert scores["double_y"]["absolute_error"] == 6. and scores["double_y"]["interval_score95"] == 164.
    assert record["last_prediction"] == snapshot and scores["y"]["comparison_kind"] == "事前预测对照"
    service.revise_record(eid, {"x": .5}, {"y": 10.}, status="completed")
    assert p.experiment(eid)["prediction_evaluation"]["rows"][0]["interval_score95"] == 2.
    assert p.experiment(eid)["last_prediction"] == snapshot
    service.undo(); assert p.experiment(eid) == record
    service.redo(); assert p.experiment(eid)["last_prediction"] == snapshot
    with Project(p.path) as reopened:
        assert reopened.experiments() == p.experiments()
        assert reopened.experiment(eid)["last_prediction"] == snapshot


def test_recommendations_save_automatic_snapshot_and_original_batch_remains_immutable(history_service):
    service = history_service; p = service.project
    batch = {"revision": p.revision, "request": {"n": 1, "model": "gp_rbf_v1", "seed": 4},
             "stage": "model_driven", "candidates": [{"conditions": {"x": .5}, "arrangement": "new",
             "prediction": {"y": {"mean": 10., "quantiles": [9.,10.,11.]}}}]}
    _, [eid] = service.commit_batch(batch)
    original = deepcopy(p.batches())
    snapshot = deepcopy(p.experiment(eid)["last_prediction"])
    service.revise_record(eid, {"x": .5}, {"y": 10.5}, status="completed")
    assert p.experiment(eid)["last_prediction"] == snapshot
    assert p.experiment(eid)["prediction_evaluation"]["rows"][0]["absolute_error"] == .5
    assert p.batches() == original and snapshot["model"] == "gp_rbf_v1"


def test_legacy_recommendation_is_recovered_before_measurement_edit(history_service):
    service = history_service; p = service.project
    _, [eid] = service.commit_batch({"revision": p.revision, "request": {"n": 1, "model": "gp_matern25_v1"},
        "candidates": [{"conditions": {"x": .5}, "prediction": {"y": {"mean": 10., "quantiles": [9.,10.,11.]}}}]})
    legacy = p.experiment(eid); legacy.pop("last_prediction"); legacy.pop("prediction_evaluation"); legacy.pop("id")
    with p.transaction(): p.db.execute("UPDATE experiments SET payload=? WHERE id=?", (encode(legacy), eid))
    report = prediction_score_report(p)
    assert report["entries"][0]["snapshot"]["source"] == "legacy_recommendation"
    service.revise_record(eid, {"x": .5}, {"y": 12.}, status="completed")
    record = p.experiment(eid)
    assert record["last_prediction"]["source"] == "legacy_recommendation"
    assert record["last_prediction"]["template_snapshot"] == p.template.data
    assert record["prediction_evaluation"]["rows"][0]["interval_score95"] == 42.


def test_explicit_reprediction_undo_and_measured_at_prediction_label(history_service):
    service = history_service; p = service.project
    eid = service.add_record({"x": .5}, {"y": 12.}, performed_at="2026-01-01")
    before = p.experiment(eid)
    service.save_record_predictions(report_for(service, [eid], mean=10.))
    first = p.experiment(eid)
    assert first["performed_at"] == before["performed_at"]
    assert all(r["comparison_kind"] == "预测时已有结果对照" for r in first["prediction_evaluation"]["rows"])
    service.save_record_predictions(report_for(service, [eid], mean=11.))
    assert p.experiment(eid)["last_prediction"]["id"] != first["last_prediction"]["id"]
    service.undo(); assert p.experiment(eid) == first
    service.undo(); assert p.experiment(eid) == before
    service.redo(); assert p.experiment(eid) == first
    assert [r["entity"] for r in p.history()].count("record_prediction") == 2


@pytest.mark.parametrize("problem", ["stale", "identity", "conditions", "length", "duplicate", "empty"])
def test_recording_stale_or_invalid_prediction_is_atomic(history_service, problem):
    service = history_service; p = service.project
    eid = add_pending(service); report = report_for(service, [eid])
    if problem == "stale": add_pending(service, .1)
    elif problem == "identity": report["source_experiments"][0]["record_uid"] = "wrong"
    elif problem == "conditions": report["conditions"][0] = {"x": .7}
    elif problem == "length": report["source_experiments"] = []
    elif problem == "duplicate":
        for key in ("conditions", "predictions", "source_experiments"): report[key] *= 2
    else: report["predictions"] = [{}]
    before = p.snapshot(), p.history()
    with pytest.raises((ValidationError, StaleVersionError)): service.save_record_predictions(report)
    assert (p.snapshot(), p.history()) == before


def test_prediction_write_failure_rolls_back_all_rows_and_revision(history_service, monkeypatch):
    service = history_service; p = service.project; ids = [add_pending(service), add_pending(service, .8)]
    report = report_for(service, ids); before = p.snapshot(), p.history()
    original = p.log
    def fail(entity, eid, *args):
        if entity == "record_prediction" and eid == ids[1]: raise RuntimeError("模拟磁盘写入失败")
        original(entity, eid, *args)
    monkeypatch.setattr(p, "log", fail)
    with pytest.raises(RuntimeError): service.save_record_predictions(report)
    assert (p.snapshot(), p.history()) == before


def test_changed_conditions_and_formula_do_not_mix_old_prediction(history_service):
    service = history_service; p = service.project; eid = add_pending(service)
    service.save_record_predictions(report_for(service, [eid]))
    frozen = deepcopy(p.experiment(eid)["last_prediction"])
    service.revise_record(eid, {"x": .6}, {"y": 10.}, status="completed")
    assert all(r["absolute_error"] is None for r in p.experiment(eid)["prediction_evaluation"]["rows"])
    service.revise_record(eid, {"x": .5}, {"y": 10.}, status="completed")
    service.update_template(p.template.revised(derived_metrics=[{"name": "double_y", "expression": "y * 3"}]))
    rows = {r["metric"]: r for r in p.experiment(eid)["prediction_evaluation"]["rows"]}
    assert rows["y"]["absolute_error"] == 0. and rows["double_y"]["absolute_error"] is None
    assert p.experiment(eid)["last_prediction"] == frozen


@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_bulk_backfill_preserves_prediction_and_rejects_forged_score_columns(history_service, tmp_path, extension):
    service = history_service; p = service.project; eid = add_pending(service)
    service.save_record_predictions(report_for(service, [eid])); snapshot = deepcopy(p.experiment(eid)["last_prediction"])
    path = tmp_path / f"backfill.{extension}"; export_records(p, path)
    assert preview_import(service, path, mode="update").commit(service) == []
    if extension == "xlsx":
        wb = load_workbook(path); columns = {c.value:c.column for c in wb.active[1]}
        for name, value in {"y": 13., "status": "completed", "last_prediction_json": "{}", "prediction_evaluation_json": "{}",
                            "上次预测评分:y:预测中心": 13., "上次预测评分:y:IS95": 0.}.items():
            wb.active.cell(2, columns[name], value)
        wb.save(path); wb.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as source: rows = list(csv.DictReader(source))
        headers=list(rows[0]); rows[0].update(y="13", status="completed", last_prediction_json="{}", prediction_evaluation_json="{}")
        with path.open("w", encoding="utf-8-sig", newline="") as target:
            writer=csv.DictWriter(target, headers); writer.writeheader(); writer.writerows(rows)
    preview = preview_import(service, path, mode="update"); assert not preview.errors
    assert preview.commit(service) == [eid]
    assert p.experiment(eid)["last_prediction"] == snapshot
    assert p.experiment(eid)["prediction_evaluation"]["rows"][0]["interval_score95"] == 82.
    service.undo(); assert p.experiment(eid)["observations"] == {}
    service.redo(); assert p.experiment(eid)["last_prediction"] == snapshot


@pytest.mark.parametrize("model", ["gp_rbf_v1", "gp_matern25_v1", "bayesian_linear_v1"])
def test_real_models_prediction_snapshot_survives_new_measurements_without_refit(history_service, model, monkeypatch):
    service = history_service; p = service.project
    for x in (0., .3, .7, 1.): service.add_record({"x": x}, {"y": {"value": 20*x, "uncertainty": {"kind": "std", "amount": .5}}})
    eid = add_pending(service)
    report = analyze(p.snapshot(), {"conditions": [{"x": .5}], "source_experiments": [{"experiment_id": eid, "record_uid": p.experiment(eid)["record_uid"]}], "model": model, "seed": 8})
    expected = deepcopy(report["predictions"][0])
    service.save_record_predictions(report)
    from experiment_planner.engine import models
    monkeypatch.setattr(models, "fit_models", lambda *a, **k: pytest.fail("Scoring must use saved values"))
    service.revise_record(eid, {"x": .5}, {"y": 12.}, status="completed")
    record = p.experiment(eid)
    assert record["last_prediction"]["prediction"] == expected
    row = next(r for r in record["prediction_evaluation"]["rows"] if r["metric"] == "y")
    assert row["absolute_error"] == pytest.approx(abs(expected["y"]["mean"]-12))
    assert row["lower95"] == expected["y"]["quantiles"][0]


def test_scores_export_contains_formulas_precise_values_and_literal_text(history_service, tmp_path):
    service = history_service; p = service.project; eid = add_pending(service)
    service.save_record_predictions(report_for(service, [eid])); service.revise_record(eid, {"x": .5}, {"y": 13.}, status="completed", experiment_batch="=1+1")
    before = p.snapshot(); report = prediction_score_report(p)
    path=tmp_path/"scores.xlsx"; export_prediction_scores(report, path)
    wb=load_workbook(path)
    assert wb.sheetnames == ["评分明细", "指标汇总", "上次预测信息", "预测条件核对", "评分公式"]
    values=list(wb.active.values); row=dict(zip(values[0],values[1]))
    assert row["上次预测中心"] == 10. and row["测量对照值"] == 13. and row["区间得分 IS95（越小越好）"] == 82.
    assert any("40×max" in str(c.value) for cells in wb["评分公式"] for c in cells)
    assert all(c.data_type != "f" for sheet in wb for cells in sheet for c in cells)
    wb.close(); assert p.snapshot() == before
    export_prediction_scores(report, tmp_path/"scores.csv")
    assert "区间得分 IS95" in (tmp_path/"scores.csv").read_text(encoding="utf-8-sig")


def test_desktop_explicit_save_and_reopen_preserves_prediction_and_score(history_service, tmp_path):
    destination=tmp_path/"session.sqlite"
    session=ProjectSession(destination, template=history_service.project.template, creating=True)
    service=PlannerService(session); eid=add_pending(service)
    service.save_record_predictions(report_for(service,[eid])); service.revise_record(eid,{"x":.5},{"y":13.},status="completed")
    expected=deepcopy(session.experiment(eid)); assert session.dirty
    session.save(); session.close()
    with Project(destination) as reopened:
        assert reopened.experiment(eid)==expected


def test_cli_predict_record_and_score_export(history_service, capsys, tmp_path):
    from experiment_planner.cli import main
    service=history_service; p=service.project
    for x in (0.,1.): service.add_record({"x":x},{"y":20*x})
    eid=add_pending(service)
    assert main(["predict-records",str(p.path),"--ids",str(eid),"--model","bayesian_linear_v1","--commit"]) == 0
    assert json.loads(capsys.readouterr().out)["recorded_ids"] == [eid]
    assert main(["scores",str(p.path),"--ids",str(eid),"--output",str(tmp_path/"cli.xlsx")]) == 0
    assert json.loads(capsys.readouterr().out)["entries"][0]["snapshot"]["id"]
    assert (tmp_path/"cli.xlsx").is_file()


def test_saving_prediction_does_not_break_external_id_import_dedup(history_service):
    service=history_service; p=service.project
    item={"external_id":"external-pending", "conditions":{"x":.5}, "observations":{}, "status":"pending"}
    [eid]=service.import_records([item],p.revision)
    service.save_record_predictions(report_for(service,[eid]))
    before=p.snapshot()
    assert service.import_records([item],p.revision)==[] and p.snapshot()==before
