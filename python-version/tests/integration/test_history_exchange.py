import copy
import csv
import json

import pytest
from openpyxl import load_workbook

from experiment_planner.domain.errors import StaleVersionError, ValidationError
from experiment_planner.domain.template import Template
from experiment_planner.io.exchange import export_import_template, export_analysis, preview_import
from experiment_planner.storage.project import Project


def test_delete_restore_and_reuse_preserve_identity_and_defaults(service, conditions, measurements):
    p = service.project
    ids = [service.add_record(conditions, measurements, note=f"记录{i}") for i in range(3)]
    service.update_template(p.template.revised(batch_defaults={**p.template.data["batch_defaults"],
        "baseline_experiment_id": 2, "repeat_ids": [2]}))
    before = copy.deepcopy(p.experiments())
    revision = p.revision
    service.delete_records([2, 3])
    service.undo()
    assert p.experiments() == before
    assert p.template.data["batch_defaults"]["baseline_experiment_id"] == 2
    assert p.template.data["batch_defaults"]["repeat_ids"] == [2]
    assert p.revision == revision + 2
    service.redo()
    eid = service.add_record(conditions, measurements, note="新的记录")
    assert eid == 2
    assert p.experiment(eid)["record_uid"] != before[1]["record_uid"]
    with pytest.raises(ValidationError, match="没有可重做"): service.redo()
    service.undo()  # undo new insertion, then undo deletion
    service.undo()
    assert p.experiments() == before
    assert [e["entity"] for e in p.history()].count("experiment_delete") == 2
    assert service.commit_batch({"revision": p.revision, "request": {"n": 1},
        "candidates": [{"conditions": conditions}]})[1] == [4]
    with pytest.raises(ValidationError, match="没有可撤销"): service.undo()


def test_undo_revise_template_and_import_with_dedup(service, conditions, measurements):
    p = service.project
    row = {"external_id": "a", "conditions": conditions, "observations": measurements}
    [eid] = service.import_records([row], p.revision)
    original = p.experiment(eid)
    service.undo()
    assert not p.experiments()
    service.redo()
    assert p.experiment(eid) == original
    assert service.import_records([row], p.revision) == []
    service.revise_record(eid, conditions, {**measurements, "sio2_remaining_nm": 80.}, note="更正")
    assert p.experiment(eid)["record_uid"] == original["record_uid"]
    service.undo()
    assert p.experiment(eid) == original
    metrics = copy.deepcopy(p.template.metrics)
    metrics[0]["expression"] = "abs(sio2_initial_nm - sio2_remaining_nm)"
    metrics[0]["display_digits"] = 2
    service.update_template(p.template.revised(derived_metrics=metrics))
    changed = p.template.data["template_version"]
    service.undo()
    assert p.template.metrics[0].get("display_digits") is None
    service.redo()
    assert p.template.metrics[0]["display_digits"] == 2
    assert p.template.data["template_version"] > changed
    with Project(p.path) as reopened:
        assert reopened.experiments() == p.experiments()
        assert reopened.history() == p.history()


def test_external_changes_cannot_be_overwritten_by_history(service, conditions, measurements):
    from experiment_planner.application.service import PlannerService
    p = service.project
    service.add_record(conditions, measurements)
    with Project(p.path) as other:
        PlannerService(other).add_record(conditions, measurements)
    with pytest.raises(StaleVersionError): service.undo()
    assert len(p.experiments()) == 2
    service.add_record(conditions, measurements)
    service.undo()
    assert len(p.experiments()) == 2
    with pytest.raises(ValidationError): service.undo()


def test_failed_undo_is_atomic(service, conditions, measurements, monkeypatch):
    p = service.project
    service.add_record(conditions, measurements)
    service.delete_records([1])
    before = p.snapshot(), p.history()
    log = p.log
    def fail(*args): raise RuntimeError("撤销写入失败")
    monkeypatch.setattr(p, "log", fail)
    with pytest.raises(RuntimeError): service.undo()
    assert (p.snapshot(), p.history()) == before
    monkeypatch.setattr(p, "log", log)
    service.undo()
    assert len(p.experiments()) == 1


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_current_template_empty_import_file_can_be_filled(service, conditions, measurements, tmp_path, extension):
    path = tmp_path / f"空表.{extension}"
    export_import_template(service.project.template, path)
    if extension == "xlsx":
        wb = load_workbook(path)
        assert wb.active.title == "导入数据" and "填写说明" in wb.sheetnames
        headers = [cell.value for cell in wb.active[1]]
        wb.active.append([({"external_id": "external-1", **conditions, **measurements}).get(key) for key in headers])
        wb.save(path); wb.close()
    else:
        with path.open(encoding="utf-8-sig") as source: headers = next(csv.reader(source))
        with path.open("a", encoding="utf-8-sig", newline="") as target:
            csv.DictWriter(target, headers).writerow({"external_id": "external-1", **conditions, **measurements})
    assert "sio2_loss_nm" not in headers
    preview = preview_import(service, path)
    assert not preview.errors
    assert len(preview.commit(service)) == 1
    assert service.project.experiment(1)["derived"]["sio2_loss_nm"]["value"] == 10.


def test_export_cached_plot_numbers_and_literal_text(template, conditions, tmp_path):
    value = 1.234567890123
    summary = {"sio2_loss_nm": {"mean": value, "quantiles": [value-.1, value, value+.1],
        "interval_kind": "latent_response"}, "raw_selectivity": {"mean": None, "quantiles": [1., 2., 3.]},
        "feasibility_probability": .75, "seed": 5}
    report = {"revision": 42, "template_version": 2, "model": "gp_rbf_v1",
        "conditions": [conditions], "predictions": [summary],
        "plot": {"conditions": [{**conditions, "etch_time_s": 20}, {**conditions, "etch_time_s": 30}],
            "predictions": [summary, summary]}}
    modified = copy.deepcopy(template.data)
    modified["derived_metrics"][0].update(display_digits=0, label="=1+1")
    template = Template(modified)
    export_analysis(report, template, tmp_path / "曲线.xlsx")
    wb = load_workbook(tmp_path / "曲线.xlsx")
    rows = list(wb.active.values)
    data = [dict(zip(rows[0], row)) for row in rows[1:]]
    assert len(data) == 6
    assert data[0]["中位数50%"] == pytest.approx(value, abs=1e-14)
    assert data[1]["后验均值"] is None
    assert [r["etch_time_s"] for r in data[2:]] == [20, 20, 30, 30]
    assert all(r["数据版本"] == 42 for r in data)
    label_column = rows[0].index("指标名称") + 1
    assert wb.active.cell(2, label_column).data_type == "s"
    wb.close()
