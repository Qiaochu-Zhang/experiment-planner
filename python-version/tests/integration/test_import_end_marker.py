import csv
import math

import pytest
from openpyxl import Workbook

from experiment_planner.application.service import PlannerService
from experiment_planner.domain.errors import ValidationError
from experiment_planner.domain.template import Template
from experiment_planner.io.exchange import preview_import
from experiment_planner.storage.project import Project


def write_table(path, headers, rows):
    if path.suffix == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as target:
            writer = csv.writer(target); writer.writerow(headers); writer.writerows(rows)
    else:
        workbook = Workbook(); sheet = workbook.active
        sheet.append(headers)
        for row in rows: sheet.append(row)
        workbook.save(path); workbook.close()


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_first_three_of_eight_stop_before_line_even_with_later_data(service, conditions, measurements, tmp_path, extension):
    names = [p['name'] for p in service.project.template.parameters]
    headers = ['external_id', *reversed(names), *measurements]
    def row(values, external): return [external, *[values.get(n) for n in reversed(names)], *measurements.values()]
    stop = {**conditions, **{n: '  ' for n in names[:3]}}
    path = tmp_path / ('records.' + extension)
    write_table(path, headers, [row(conditions, 'one'), row(stop, 'marker-with-measurements'), row(conditions, 'must-not-import')])
    preview = preview_import(service, path)
    assert preview.errors == [] and len(preview.records) == 1
    assert preview.stop_row == 3 and preview.stop_fields == names[:3]
    assert '后续行不导入' in preview.stop_message
    assert len(preview.commit(service)) == 1
    assert len(service.project.experiments()) == 1


@pytest.mark.parametrize('count', [1, 3, 4, 8, 10, 11])
def test_end_count_rounds_up_and_excludes_measurements(tmp_path, count):
    names = [f'x{i}' for i in range(count)]
    template = Template({'schema_version': 2, 'template_version': 1, 'parameters': [{'name': n, 'bounds': [0, 10]} for n in names],
                         'measurements': [{'name': 'y', 'is_response': True}], 'objectives': [{'metric': 'y', 'direction': 'maximize'}],
                         'optimization': {'mode': 'single'}, 'ratio_policy': {'mode': 'remove'}})
    with Project.create(tmp_path / 'custom.sqlite', '任意输入个数', template) as project:
        service = PlannerService(project); k = math.ceil(count * .3)
        path = tmp_path / 'custom.csv'
        write_table(path, [*names, 'y'], [[0] * count + [1], [None] * k + [2] * (count - k) + [9], [3] * count + [8]])
        preview = preview_import(service, path)
        assert preview.stop_row == 3 and preview.stop_fields == names[:k]
        assert len(preview.records) == 1 and not preview.errors


def test_missing_headers_do_not_masquerade_as_end(service, tmp_path):
    path = tmp_path / 'wrong.csv'; path.write_text('unknown\n\n', encoding='utf-8')
    preview = preview_import(service, path)
    assert preview.errors[0]['row'] == 1 and preview.stop_row is None
    with pytest.raises(ValidationError): preview.commit(service)


def test_partial_missing_before_end_still_blocks_commit(service, conditions, tmp_path):
    names = [p['name'] for p in service.project.template.parameters]
    values = [conditions[n] for n in names]; values[0] = None
    path = tmp_path / 'partial.csv'
    write_table(path, names, [values, [None] * len(names), [conditions[n] for n in names]])
    preview = preview_import(service, path)
    assert preview.stop_row == 3 and preview.errors[0]['row'] == 2
    with pytest.raises(ValidationError): preview.commit(service)
    assert service.project.experiments() == []


def test_csv_completely_empty_row_and_mapped_headers(service, conditions, tmp_path):
    names = [p['name'] for p in service.project.template.parameters]
    headers = [f'列{i}' for i in range(len(names))]
    path = tmp_path / 'blank.csv'
    write_table(path, headers, [[conditions[n] for n in names], [], [conditions[n] for n in names]])
    preview = preview_import(service, path, dict(zip(headers, names)))
    assert preview.stop_row == 3 and len(preview.records) == 1 and not preview.errors


def test_excel_formatted_tail_and_eof(service, conditions, tmp_path):
    names = [p['name'] for p in service.project.template.parameters]
    path = tmp_path / 'styled.xlsx'
    workbook = Workbook(); sheet = workbook.active
    sheet.append(names); sheet.append([conditions[n] for n in names])
    sheet.cell(24, 1).number_format = '0.000'
    workbook.save(path); workbook.close()
    preview = preview_import(service, path)
    assert preview.stop_row == 3 and len(preview.records) == 1 and not preview.errors
    path = tmp_path / 'end.csv'
    write_table(path, names, [[conditions[n] for n in names]])
    preview = preview_import(service, path)
    assert preview.stop_row is None and len(preview.records) == 1


def test_false_and_zero_are_data_not_end(tmp_path):
    template = Template({'schema_version': 2, 'template_version': 1,
        'parameters': [{'name': 'enabled', 'value_type': 'bool'}],
        'measurements': [{'name': 'score', 'is_response': True}], 'objectives': [{'metric': 'score', 'direction': 'maximize'}],
        'optimization': {'mode': 'single'}, 'ratio_policy': {'mode': 'remove'}})
    with Project.create(tmp_path / 'boolean.sqlite', '是否输入', template) as project:
        path = tmp_path / 'boolean.xlsx'; write_table(path, ['enabled', 'score'], [[False, 0], [None, 2], [True, 1]])
        preview = preview_import(PlannerService(project), path)
        assert not preview.errors and preview.stop_row == 3 and preview.records[0]['conditions'] == {'enabled': False}
