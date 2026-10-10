import copy

import pytest
from openpyxl import load_workbook

from experiment_planner.application.service import PlannerService
from experiment_planner.application.template_editing import blank_template
from experiment_planner.domain.errors import ValidationError
from experiment_planner.domain.template import Template
from experiment_planner.engine.models import Encoder, fit_models
from experiment_planner.io.exchange import export_import_template, preview_import, read_headers
from experiment_planner.io.template_excel import read_template_excel, write_template_excel
from experiment_planner.storage.project import Project


def varied_template():
    data = blank_template()
    data['parameters'] = [
        {'name': 'speed', 'unit': 's', 'bounds': [0, 10], 'execution_rounding': {'mode': 'digits', 'digits': 0}},
        {'name': 'count', 'value_type': 'int', 'bounds': [1, 4]},
        {'name': 'gas', 'value_type': 'category', 'values': ['混合,气体', '甲' * 280, '=literal']},
        {'name': 'switch', 'value_type': 'bool'},
        {'name': 'sample', 'value_type': 'text', 'values': ['批次甲', '批次乙']},
        {'name': 'constant', 'fixed_value': 3.5},
        {'name': 'level', 'value_type': 'int', 'values': [1, 2, 4]},
    ]
    data['measurements'] = [
        {'name': 'quality', 'unit': 'kg', 'is_response': True, 'model_base': True},
        {'name': 'grade', 'value_type': 'category', 'values': ['合格', '待复测']},
        {'name': 'approved', 'value_type': 'bool'},
    ]
    data['objectives'] = [{'metric': 'quality', 'direction': 'maximize'}]
    return Template(data)


def validation_at(sheet, cell):
    return next(v for v in sheet.data_validations.dataValidation if cell in v)


def options_at(workbook, sheet, cell):
    validation = validation_at(sheet, cell)
    assert validation.type == 'list' and validation.showDropDown is False
    assert validation.showErrorMessage and validation.showInputMessage and validation.errorStyle == 'stop'
    defined = workbook.defined_names[validation.formula1.removeprefix('=')]
    sheet_name, address = next(defined.destinations)
    assert workbook[sheet_name].sheet_state == 'hidden'
    return [cell.value for row in workbook[sheet_name][address] for cell in row]


def append_record(sheet, values):
    sheet.append([values.get(cell.value) for cell in sheet[2]])


def record_values():
    return {'external_id': 'test', 'speed': .25, 'count': 2, 'gas': '混合,气体',
            'switch': '是', 'sample': '批次甲', 'constant': 3.5, 'level': 2,
            'quality': 5, 'grade': '合格', 'approved': 'false'}


def test_exported_record_template_has_required_range_hints_and_reusable_dropdowns(tmp_path):
    template = varied_template(); path = tmp_path / '记录.xlsx'
    export_import_template(template, path)
    workbook = load_workbook(path); sheet = workbook.active
    try:
        headers = {c.value: c.column_letter for c in sheet[2]}
        assert read_headers(path) == [c.value for c in sheet[2]]
        assert sheet.max_row == 2 and sheet.freeze_panes == 'A3'
        speed = headers['speed']
        assert '必填' in sheet[f'{speed}1'].value and '0 ≤ 值 ≤ 10' in sheet[f'{speed}1'].value
        assert '单位 s' in sheet[f'{speed}2'].comment.text
        assert '>=' in validation_at(sheet, f'{speed}3').formula1
        count_rule = validation_at(sheet, f"{headers['count']}3")
        assert 'INT(' in count_rule.formula1 and '>=1' in count_rule.formula1 and '<=4' in count_rule.formula1
        assert options_at(workbook, sheet, f"{headers['gas']}3") == template.parameters[2]['values']
        assert options_at(workbook, sheet, f"{headers['gas']}10002") == template.parameters[2]['values']
        assert options_at(workbook, sheet, f"{headers['constant']}3") == ['3.5']
        assert options_at(workbook, sheet, f"{headers['level']}3") == ['1', '2', '4']
        assert options_at(workbook, sheet, f"{headers['sample']}3") == ['批次甲', '批次乙']
        assert 'true' in options_at(workbook, sheet, f"{headers['switch']}3")
        assert options_at(workbook, sheet, f"{headers['grade']}3") == ['合格', '待复测']
        assert 'false' in options_at(workbook, sheet, f"{headers['approved']}3")
        assert options_at(workbook, sheet, f"{headers['approved_uncertainty_kind']}3") == ['unspecified']
        assert options_at(workbook, sheet, f"{headers['grade_uncertainty_kind']}3") == ['unspecified']
        assert options_at(workbook, sheet, f"{headers['quality_unit']}3") == ['kg']
        assert 'interval' in options_at(workbook, sheet, f"{headers['quality_uncertainty_kind']}3")
        assert options_at(workbook, sheet, f"{headers['quality_distribution']}3") == ['normal', 'uniform']
        assert '>0' in validation_at(sheet, f"{headers['quality_coverage_factor']}3").formula1
        confidence = headers['quality_confidence']
        assert '0 < 数值 < 1' in sheet[f'{confidence}1'].value
        assert '<1' in validation_at(sheet, f'{confidence}3').formula1
        append_record(sheet, {**record_values(), 'quality_uncertainty': .5, 'quality_uncertainty_kind': 'interval',
                              'quality_distribution': 'normal', 'quality_confidence': .95})
        workbook.save(path)
    finally: workbook.close()
    with Project.create(tmp_path / 'records.sqlite', '测试', template) as project:
        service = PlannerService(project); preview = preview_import(service, path)
        assert not preview.errors and len(preview.records) == 1
        preview.commit(service)
        assert project.experiment(1)['actual']['speed'] == .25  # actuals need not obey recommendation grid
        assert project.experiment(1)['actual']['switch'] is True
        assert project.experiment(1)['observations']['approved']['value'] is False
        assert project.experiment(1)['observations']['quality']['uncertainty']['confidence'] == .95


@pytest.mark.parametrize('updates,message', [
    ({'speed': 11}, '超出范围'), ({'count': 1.5}, '整数'), ({'gas': '未知气体'}, '未知类别'),
    ({'constant': 4}, '固定值'), ({'level': 3}, '离散值'), ({'count': None}, '缺少输入'),
    ({'quality_confidence': 1}, '置信水平'), ({'quality_uncertainty': -1}, '非负'),
])
def test_excel_paste_cannot_bypass_software_validation(tmp_path, updates, message):
    template = varied_template(); path = tmp_path / '粘贴.xlsx'
    export_import_template(template, path)
    workbook = load_workbook(path)
    append_record(workbook.active, {**record_values(), **updates})
    workbook.save(path); workbook.close()
    with Project.create(tmp_path / 'records.sqlite', '测试', template) as project:
        service = PlannerService(project); preview = preview_import(service, path)
        assert preview.errors[0]['row'] == 3 and message in preview.errors[0]['message']
        with pytest.raises(ValidationError): preview.commit(service)
        assert not project.experiments()


def test_guided_import_keeps_physical_stop_and_header_error_rows(tmp_path, service, conditions):
    path = tmp_path / '结束.xlsx'; export_import_template(service.project.template, path)
    workbook = load_workbook(path); sheet = workbook.active
    append_record(sheet, conditions); append_record(sheet, {'note': '结束'})
    append_record(sheet, {**conditions, 'cl2_sccm': '坏数据'})
    workbook.save(path); workbook.close()
    preview = preview_import(service, path)
    assert not preview.errors and len(preview.records) == 1 and preview.stop_row == 4
    mapping = {'cl2_sccm': 'elsewhere'}
    preview = preview_import(service, path, mapping)
    assert preview.errors[0]['row'] == 2 and '未匹配' in preview.errors[0]['message']


def test_definition_dropdowns_settings_notes_and_editable_units_roundtrip(tmp_path, template):
    data = template.data; data['custom_bool'] = True
    path = tmp_path / '定义.xlsx'; write_template_excel(data, path)
    workbook = load_workbook(path)
    try:
        parameters = workbook['输入参数']; assert parameters['A2'].value == '内部名称'
        assert '必填' in parameters['A1'].value and '-12' in parameters['O1'].value
        assert '文字' in options_at(workbook, parameters, 'C3')
        assert '是' in options_at(workbook, parameters, 'J3')
        assert '指定步长' in options_at(workbook, parameters, 'P3')
        assert options_at(workbook, workbook['优化目标'], 'B3') == ['maximize', 'minimize', '尽量大', '尽量小']
        assert options_at(workbook, workbook['输入约束'], 'B3') == ['<=', '>=']
        assert 'monotonic' in options_at(workbook, workbook['先验关系'], 'D3')
        assert '必填' in workbook['派生指标与公式']['N1'].value
        settings = workbook['模板设置']
        rows = {settings.cell(r, 1).value: r for r in range(3, settings.max_row + 1)}
        base_row = rows['modeling.base_responses']
        assert '说明性记录' in settings.cell(base_row, 4).value
        assert '不自动排除' in settings.cell(rows['context_assumptions'], 4).value
        model_row = rows['optimization.default_model_preset']
        assert 'gp_matern25_v1' in options_at(workbook, settings, f'B{model_row}')
        settings.cell(model_row, 2, 'gp_matern25_v1')
        settings.cell(rows['custom_bool'], 2, '否')
        assert options_at(workbook, settings, 'C10002') == ['str', 'int', 'float', 'bool', 'list', 'dict', 'null']
        # Unit is free text: there must be no enum restricting this cell.
        assert not any('D3' in v for v in parameters.data_validations.dataValidation)
        parameters['D3'] = '自定义流量'
        workbook.save(path)
    finally: workbook.close()
    restored = Template.read(path)
    assert restored.parameters[0]['unit'] == '自定义流量'
    assert restored.data['optimization']['default_model_preset'] == 'gp_matern25_v1'
    assert restored.data['custom_bool'] is False
    assert restored.data['modeling'] == data['modeling']
    assert restored.data['context_assumptions'] == data['context_assumptions']


def test_blank_definition_can_be_filled_from_row_three(tmp_path):
    path = tmp_path / '空表.xlsx'; write_template_excel(blank_template(), path)
    workbook = load_workbook(path)
    for title, values in [
        ('输入参数', ['speed', '速度', '数值', '1', None, 0, 10]),
        ('原始测量', ['quality', '质量', '数值', 'kg']),
        ('派生指标与公式', ['double_quality', '双倍质量', 'float', 'kg', None, None, None, None, None, None, '是', '否', None, 'quality * 2']),
        ('优化目标', ['double_quality', '尽量大']),
    ]:
        workbook[title].append(values)
    workbook.save(path); workbook.close()
    template = Template.read(path)
    assert template.graph.evaluate({'quality': 3})['double_quality'].value == 6
    assert [f['name'] for f in template.responses] == ['quality']
    assert read_template_excel(path)['parameters'][0]['bounds'] == [0, 10]


def test_old_experiment_metadata_does_not_select_models_or_inputs():
    data = blank_template()
    data['parameters'] = [{'name': 'speed', 'bounds': [0, 10]}]
    data['measurements'] = [{'name': 'quality', 'is_response': True, 'model_base': True}]
    data['objectives'] = [{'metric': 'quality', 'direction': 'maximize'}]
    clean = Template(data)
    old = copy.deepcopy(data)
    old['modeling'] = {'base_responses': ['sio2_loss_nm', 'sin_loss_nm']}
    old['context_assumptions'] = [{'field': 'speed', 'response': 'quality', 'assumption': 'not_used_as_response_model_input'}]
    stale = Template(old)
    records = [{'actual': {'speed': x}, 'status': 'completed', 'observations': {'quality': {'value': 2 * x + 1, 'uncertainty': {}}}, 'derived': {}} for x in (1, 3, 6)]
    for template in (clean, stale):
        assert [f['name'] for f in template.responses] == ['quality']
        assert Encoder(template).width == 1
        _, encoder, names, _, _ = fit_models(template, records, 'bayesian_linear_v1')
        assert names == ['quality'] and encoder.width == 1
