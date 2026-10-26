import copy
import csv
import time

import pytest
import torch
from openpyxl import load_workbook

from experiment_planner.application.importance import analyze_importance
from experiment_planner.application.service import PlannerService
from experiment_planner.domain.errors import ValidationError, StaleVersionError, CapabilityError
from experiment_planner.domain.reliability import data_weights, experiment_time
from experiment_planner.domain.template import Template
from experiment_planner.engine.models import fit_models
from experiment_planner.io.exchange import export_records, preview_import, export_importance
from experiment_planner.storage.project import Project, encode


@pytest.fixture
def quality_service(tmp_path):
    template = Template({"schema_version": 2, "template_version": 1, "name": "通用合成实验",
        "parameters": [{"name": "x", "bounds": [0., 1.]}, {"name": "z", "bounds": [0., 1.]},
                       {"name": "sample", "value_type": "text", "fixed_value": "样品"}],
        "measurements": [{"name": "y", "unit": "1", "is_response": True}],
        "derived_metrics": [{"name": "double_y", "expression": "y * 2", "unit": "1", "model_base": False}],
        "objectives": [{"metric": "y", "direction": "maximize"}], "optimization": {"mode": "single"},
        "ratio_policy": {"mode": "remove"}})
    project = Project.create(tmp_path / "quality.sqlite", "可信度合成验证", template)
    yield PlannerService(project)
    project.close()


def add(quality_service, x=.5, z=.5, y=1., **kwargs):
    return quality_service.add_record({"x": x, "z": z, "sample": "样品"},
        {"y": {"value": y, "uncertainty": {"kind": "std", "amount": 1.}}}, **kwargs)


def edit_table(path, changes):
    wb = load_workbook(path)
    sheet = wb.active; headers = {cell.value: cell.column for cell in sheet[1]}
    for row, values in changes.items():
        for key, value in values.items(): sheet.cell(row, headers[key]).value = value
    wb.save(path); wb.close()


def test_batch_assignment_confidence_dates_undo_reopen_and_prediction_identity(quality_service):
    service = quality_service; p = service.project
    a, b = add(service, performed_at="2026-01-01"), add(service, performed_at="2026-02-01")
    before = copy.deepcopy(p.experiments())
    service.assign_experiment_batch([a, b], " 校准批次 ", confidence=.4)
    assert all(e["experiment_batch"] == "校准批次" for e in p.experiments())
    assert [e["performed_at"] for e in p.experiments()] == [e["performed_at"] for e in before]
    assert p.template.data["data_weighting"]["batch_confidence"] == {"校准批次": .4}
    service.undo(); assert p.experiments() == before
    assert not p.template.data.get("data_weighting")
    service.redo()
    service.assign_experiment_batch([a], "校准批次", performed_at="2026-03-01")
    with Project(p.path) as reopened:
        assert reopened.experiments() == p.experiments()
        assert reopened.experiment(b)["performed_at"] == before[1]["performed_at"]
        assert reopened.experiment(a)["record_uid"] == before[0]["record_uid"]


@pytest.mark.parametrize("confidence", [0, -1, 1.1, float("nan"), True])
def test_invalid_confidence_batch_write_is_atomic(quality_service, confidence):
    service = quality_service; eid = add(service)
    before = service.project.snapshot(), service.project.history()
    with pytest.raises(ValidationError): service.assign_experiment_batch([eid], "批次", confidence=confidence)
    assert (service.project.snapshot(), service.project.history()) == before


def test_recency_combines_confidence_and_stable_experiment_time(quality_service):
    service = quality_service; p = service.project
    a = add(service, performed_at="2026-01-01", experiment_batch="旧")
    b = add(service, performed_at="2026-01-31")
    service.update_template(p.template.revised(data_weighting={"recency_half_life_days": 30, "batch_confidence": {"旧": .4}}))
    weights, info = data_weights(p.template, p.experiments())
    assert weights == pytest.approx([.2, 1.])
    assert info["reference_time"].startswith("2026-01-31")
    old_time = p.experiment(a)["performed_at"]
    service.revise_record(a, p.experiment(a)["actual"], p.experiment(a)["observations"], note="修改备注")
    assert p.experiment(a)["performed_at"] == old_time
    assert data_weights(p.template, p.experiments())[0] == pytest.approx(weights)
    assert experiment_time("2026-01-01T08:00:00+08:00") == "2026-01-01T00:00:00+00:00"
    assert p.experiment(b)["experiment_batch"] == ""


def test_unknown_legacy_dates_and_extreme_decay(quality_service):
    t = quality_service.project.template.revised(data_weighting={"recency_half_life_days": .001})
    records = [{"id": 1, "performed_at": None}, {"id": 2, "updated": "2020-01-01"}, {"id": 3, "performed_at": "2026-01-01"}]
    weights, info = data_weights(t, records)
    assert weights == [1., 1e-6, 1.]
    assert info["records"][0]["unknown_time"]
    assert info["records"][1]["legacy_time"] and info["records"][1]["weight_floor_applied"]


@pytest.mark.parametrize("model", ["gp_rbf_v1", "gp_matern25_v1", "bayesian_linear_v1"])
@pytest.mark.parametrize("known", [True, False])
def test_confidence_changes_real_models_and_default_preserves_results(quality_service, model, known):
    service = quality_service; p = service.project
    add(service, y=0., experiment_batch="不稳", performed_at="2026-01-01")
    add(service, y=10., performed_at="2026-02-01")
    if not known:
        for e in p.experiments():
            service.revise_record(e["id"], e["actual"], {"y": e["observations"]["y"]["value"]})
    fitted, encoder, _, _, _ = fit_models(p.template, p.experiments(), model, 3)
    X = encoder.encode([p.experiment(1)["actual"]])
    baseline = fitted.posterior(X).mean.detach()
    equal = p.template.revised(data_weighting={"recency_half_life_days": 0, "batch_confidence": {"不稳": 1.}})
    same, _, _, _, _ = fit_models(equal, p.experiments(), model, 3)
    assert torch.equal(same.posterior(X).mean.detach(), baseline)
    weighted = p.template.revised(data_weighting={"batch_confidence": {"不稳": .01}})
    fitted, _, _, datasets, notices = fit_models(weighted, p.experiments(), model, 3)
    assert float(fitted.posterior(X).mean.detach()[0, 0]) > float(baseline[0, 0]) + 1.
    assert [r["weight"] for r in datasets["y"]["data_weighting"]["records"]] == [.01, 1.]
    assert any("已应用批次" in n for n in notices)
    if not known and model.startswith("gp_"):
        noise = fitted.models[0].inner.likelihood.noise_covar(shape=torch.Size([2])).diagonal()
        assert float((noise[0] / noise[1]).detach()) == pytest.approx(100.)
        assert torch.isfinite(fitted.posterior(encoder.encode([{"x": .2, "z": .8, "sample": "样品"}])).mean).all()


def test_permutation_importance_dominant_parameter_and_constant_or_excluded(quality_service, tmp_path):
    service = quality_service; p = service.project
    for x in (0., .33, .66, 1.):
        for z in (0., 1.): add(service, x=x, z=z, y=10*x+.1*z)
    service.add_record({"x": .5, "z": .5, "sample": "样品"}, {"y": 999.}, status="pending")
    before = p.snapshot()
    report = analyze_importance(before, {"model": "bayesian_linear_v1", "seed": 19})
    assert report["ranking"][0]["parameter"] == "x" and report["ranking"][0]["score"] > 99.
    assert next(r for r in report["ranking"] if r["parameter"] == "sample")["score"] is None
    assert sum(r["score"] or 0 for r in report["ranking"]) == pytest.approx(100.)
    assert report["responses"][0]["count"] == 8
    assert report == analyze_importance(before, {"model": "bayesian_linear_v1", "seed": 19})
    assert p.snapshot() == before
    path = tmp_path / "importance.xlsx"; export_importance(report, p.template, path)
    wb = load_workbook(path); rows = list(wb.active.values); wb.close()
    assert rows[1][2] == pytest.approx(report["ranking"][0]["score"])
    add(service, y=4.); add(service, y=5.)
    assert p.revision != report["revision"]


def test_importance_with_missing_response_and_coupled_constraints(quality_service):
    service = quality_service; p = service.project
    measurements = [*p.template.measurements, {"name": "unmeasured", "is_response": True}]
    service.update_template(p.template.revised(measurements=measurements))
    add(service, x=0., z=0., y=0.); add(service, x=1., z=1., y=10.)
    template = p.template.revised(input_constraints=[{"coefficients": {"x": 1, "z": -1}, "op": "<=", "rhs": 0},
                                                   {"coefficients": {"x": 1, "z": -1}, "op": ">=", "rhs": 0}])
    snapshot = {**p.snapshot(), "template": template.data}
    report = analyze_importance(snapshot, {"model": "bayesian_linear_v1"})
    assert report["unavailable_responses"][0]["response"] == "unmeasured"
    assert all(r["score"] is None for r in report["ranking"])
    assert "无合法条件" in report["responses"][0]["parameters"][0]["reason"]


def test_importance_requires_real_results_and_does_not_rank_fixed_inputs(quality_service):
    service = quality_service
    with pytest.raises(CapabilityError, match="两条"): analyze_importance(service.project.snapshot(), {})
    add(service, y=2.); add(service, y=3.)
    report = analyze_importance(service.project.snapshot(), {"model": "bayesian_linear_v1"})
    assert all(r["score"] is None for r in report["ranking"])


@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_export_import_update_noop_changes_recompute_and_undo(quality_service, tmp_path, extension):
    service = quality_service; p = service.project
    a = add(service, note="保留备注", performed_at="2026-01-01")
    b = service.add_record({"x": .1, "z": .2, "sample": "样品"}, {}, status="pending")
    path = tmp_path / f"records.{extension}"; export_records(p, path)
    before = p.snapshot(); preview = preview_import(service, path, mode="update")
    assert not preview.errors and all(not e["changes"] for e in preview.changes)
    assert preview.commit(service) == [] and p.snapshot() == before
    if extension == "xlsx":
        edit_table(path, {2: {"y": 7., "experiment_batch": "新批", "performed_at": "2026-01-10"},
                          3: {"y": 8., "status": "completed", "y_note": "批量回填"}})
    else:
        with path.open(encoding="utf-8-sig", newline="") as source: rows = list(csv.DictReader(source))
        headers = list(rows[0]); rows[0].update(y="7", experiment_batch="新批", performed_at="2026-01-10")
        rows[1].update(y="8", status="completed", y_note="批量回填")
        with path.open("w", encoding="utf-8-sig", newline="") as target:
            writer = csv.DictWriter(target, headers); writer.writeheader(); writer.writerows(rows)
    preview = preview_import(service, path, mode="update")
    assert not preview.errors
    assert preview.commit(service) == [a, b]
    assert len(p.experiments()) == 2 and p.experiment(a)["derived"]["double_y"]["value"] == 14.
    assert p.experiment(a)["note"] == "保留备注" and p.experiment(b)["status"] == "completed"
    assert p.experiment(a)["record_uid"] == before["experiments"][0]["record_uid"]
    service.undo(); assert p.experiments() == before["experiments"]
    service.redo(); assert p.experiment(b)["observations"]["y"]["value"] == 8.


@pytest.mark.parametrize("problem", ["revised", "reused", "project", "identity", "status", "number", "duplicate"])
def test_bulk_update_refuses_invalid_or_stale_identity(quality_service, tmp_path, problem):
    service = quality_service; p = service.project; eid = add(service)
    path = tmp_path / "records.xlsx"; export_records(p, path)
    if problem == "revised": service.revise_record(eid, p.experiment(eid)["actual"], {"y": 2.})
    elif problem == "reused": service.delete_records([eid]); add(service)
    elif problem == "project": edit_table(path, {2: {"project_uid": "other-project"}})
    elif problem == "identity": edit_table(path, {2: {"record_uid": "other-record"}})
    elif problem == "status": edit_table(path, {2: {"status": "unknown"}})
    elif problem == "number": edit_table(path, {2: {"experiment_id": 1.5}})
    else:
        wb = load_workbook(path); wb.active.append([c.value for c in wb.active[2]]); wb.save(path); wb.close()
    before = p.snapshot(), p.history()
    preview = preview_import(service, path, mode="update")
    assert preview.errors
    with pytest.raises(ValidationError): preview.commit(service)
    assert (p.snapshot(), p.history()) == before


def test_bulk_update_revision_guard_and_transaction_rollback(quality_service, tmp_path, monkeypatch):
    service = quality_service; p = service.project
    add(service); add(service, x=.1)
    path = tmp_path / "records.xlsx"; export_records(p, path); edit_table(path, {2: {"y": 10.}, 3: {"y": 20.}})
    preview = preview_import(service, path, mode="update")
    before = p.snapshot(), p.history()
    original = p.log
    def fail(entity, eid, *args):
        if entity == "bulk_update" and eid == 2: raise RuntimeError("写入中断")
        return original(entity, eid, *args)
    monkeypatch.setattr(p, "log", fail)
    with pytest.raises(RuntimeError): preview.commit(service)
    assert (p.snapshot(), p.history()) == before
    monkeypatch.setattr(p, "log", original)
    add(service)
    with pytest.raises(StaleVersionError): preview.commit(service)


def test_importance_worker_real_process(quality_service):
    from experiment_planner.worker.process import CalculationTask
    add(quality_service, x=0., y=0.); add(quality_service, x=1., y=10.)
    task = CalculationTask(quality_service.project, {"model": "bayesian_linear_v1"}, operation="importance")
    deadline = time.monotonic() + 50
    try:
        while time.monotonic() < deadline:
            result = task.poll()
            if result is not None: break
            time.sleep(.05)
        else: pytest.fail("重要度后台进程未完成")
        assert "error" not in result, result
        assert result["result"]["ranking"][0]["parameter"] == "x"
    finally:
        if task.key in CalculationTask.active: task.cancel()


@pytest.mark.parametrize("model", ["gp_rbf_v1", "gp_matern25_v1", "bayesian_linear_v1"])
def test_recency_affects_real_model_and_weighted_recommendation(quality_service, model):
    from experiment_planner.application.analysis import analyze
    from experiment_planner.engine.planner import generate
    from experiment_planner.engine.space import BatchRequest
    service = quality_service; p = service.project
    add(service, y=0., performed_at="2026-01-01")
    add(service, y=10., performed_at="2026-02-01")
    service.update_template(p.template.revised(data_weighting={"recency_half_life_days": 5.}))
    report = analyze(p.snapshot(), {"model": model, "conditions": [p.experiment(1)["actual"]]})
    assert report["predictions"][0]["y"]["mean"] > 8.
    batch = generate(p.snapshot(), BatchRequest(n=1, model=model, pool_size=8))
    assert batch["stage"] == "model_driven" and len(batch["candidates"]) == 1
    assert batch["candidates"][0]["prediction"]["y"]["mean"] is not None
    assert batch["model_datasets"]["y"]["data_weighting"]["recency_half_life_days"] == 5.


def test_weighted_bernoulli_and_category_importance(quality_service):
    service = quality_service; p = service.project
    data = copy.deepcopy(p.template.data)
    data["measurements"][0]["value_type"] = "bool"
    data["derived_metrics"] = []
    data["parameters"][1] = {"name": "z", "value_type": "category", "values": ["甲", "乙"]}
    data["template_version"] += 1
    service.update_template(Template(data))
    for z, y in (("甲", False), ("甲", False), ("乙", True), ("乙", True)):
        service.add_record({"x": .5, "z": z, "sample": "样品"}, {"y": y}, experiment_batch="可信" if y else "较弱")
    service.assign_experiment_batch([1, 2], "较弱", confidence=.1)
    model, encoder, _, datasets, _ = fit_models(p.template, p.experiments(), "bayesian_linear_v1")
    probabilities = model.posterior(encoder.encode([e["actual"] for e in p.experiments()])).rsample(torch.Size([128]))
    assert torch.isfinite(probabilities).all() and torch.all((probabilities >= 0) & (probabilities <= 1))
    assert [r["weight"] for r in datasets["y"]["data_weighting"]["records"]] == [.1, .1, 1., 1.]
    report = analyze_importance(p.snapshot(), {"model": "bayesian_linear_v1"})
    assert report["ranking"][0]["parameter"] == "z" and report["ranking"][0]["score"] == pytest.approx(100.)


def test_bulk_update_can_clear_measurement_preserve_detailed_errors_and_legacy_time(quality_service, tmp_path):
    service = quality_service; p = service.project
    eid = add(service, performed_at="2026-01-01")
    old = p.experiment(eid); old.pop("performed_at"); old["updated"] = "2025-12-01T00:00:00+00:00"
    old["observations"]["y"]["uncertainty"] = {"kind": "bounds", "amount": 2., "distribution": "uniform", "source": "共同校准"}
    payload = {k: v for k, v in old.items() if k != "id"}
    with p.transaction(): p.db.execute("UPDATE experiments SET payload=? WHERE id=?", (encode(payload), eid))
    path = tmp_path / "legacy.xlsx"; export_records(p, path)
    preview = preview_import(service, path, mode="update")
    assert not preview.errors and preview.commit(service) == []
    edit_table(path, {2: {"note": "批量更正"}})
    preview = preview_import(service, path, mode="update"); assert not preview.errors
    preview.commit(service)
    assert p.experiment(eid)["performed_at"].startswith("2025-12-01")
    assert p.experiment(eid)["observations"]["y"]["uncertainty"]["source"] == "共同校准"
    export_records(p, path); edit_table(path, {2: {"y": None}})
    preview = preview_import(service, path, mode="update"); assert not preview.errors
    preview.commit(service)
    assert p.experiment(eid)["observations"]["y"]["value"] is None
    assert p.experiment(eid)["derived"]["double_y"]["value"] is None
