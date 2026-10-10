import copy

import pytest
from openpyxl import Workbook, load_workbook

from experiment_planner.application.service import PlannerService
from experiment_planner.application.template_editing import blank_template, paste_fields
from experiment_planner.domain.errors import ValidationError, StaleVersionError
from experiment_planner.domain.template import Template
from experiment_planner.engine.ax_bridge import rebuild
from experiment_planner.engine.models import Encoder, fit_models
from experiment_planner.engine.objectives import prediction_summary
from experiment_planner.engine.space import BatchRequest, candidate_pool
from experiment_planner.io.template_excel import read_template_excel, write_template_excel
from experiment_planner.storage.project import Project
from experiment_planner.storage.session import ProjectSession


def small_template(text=True):
    data = blank_template()
    data['parameters'] = [{'name': 'speed', 'bounds': [0, 10], 'unit': '1'}]
    if text: data['parameters'].append({'name': 'sample', 'value_type': 'text', 'unit': '1', 'use_as_model_input': False})
    data['measurements'] = [{'name': 'quality', 'unit': 'kg', 'value_type': 'float', 'is_response': True, 'model_base': True, 'measured_at': 'after'}]
    data['derived_metrics'] = [{'name': 'double_quality', 'expression': 'quality * 2', 'unit': 'kg', 'value_type': 'float', 'model_base': False, 'is_response': True}]
    data['objectives'] = [{'metric': 'double_quality', 'direction': 'maximize'}]
    return Template(data)


def test_excel_export_import_keeps_formulas_types_settings_and_literal_text(tmp_path, template):
    data = template.data
    data['custom'] = {'list': [None, False, 1.25, '中文'], 'empty': {}}
    data['name'] = '123'
    data['parameters'][0]['note'] = '=HYPERLINK("test")'
    data['parameters'][0]['execution_rounding']['custom'] = '扩展'
    file = tmp_path / '模板.xlsx'
    Template(data).write(file)
    workbook = load_workbook(file)
    assert workbook['输入参数']['E2'].data_type == 's'
    workbook.close()
    restored = Template.read(file)
    assert restored.data['name'] == '123' and restored.data['custom'] == data['custom']
    assert restored.parameters[0]['note'] == data['parameters'][0]['note']
    assert restored.parameters[0]['execution_rounding']['custom'] == '扩展'
    assert restored.graph.evaluate({'sio2_initial_nm': 100, 'sio2_remaining_nm': 90, 'sin_initial_nm': 100, 'sin_remaining_nm': 98})['selectivity'].value == 5


def test_handwritten_chinese_excel_is_a_template(tmp_path):
    workbook = Workbook(); workbook.active.title = '输入参数'
    workbook.active.append(['内部名称', '类型', '单位', '下限', '上限', '离散值'])
    workbook.active.append(['speed', '数值', '1', 0, 10, None])
    workbook.active.append(['sample', '文字', '1', None, None, '甲，乙'])
    sheet = workbook.create_sheet('原始测量')
    sheet.append(['内部名称', '类型', '单位']); sheet.append(['quality', '数值', 'kg'])
    sheet = workbook.create_sheet('派生指标与公式')
    sheet.append(['内部名称', '单位', '计算公式']); sheet.append(['double_quality', 'kg', '=quality * 2'])
    sheet = workbook.create_sheet('优化目标')
    sheet.append(['响应或指标', '方向']); sheet.append(['double_quality', '尽量大'])
    sheet = workbook.create_sheet('模板设置')
    sheet.append(['属性', '值']); sheet.append(['name', '手写模板']); sheet.append(['optimization.mode', '单目标'])
    file = tmp_path / '手写.xlsx'; workbook.save(file); workbook.close()
    template = Template.read(file)
    assert template.data['name'] == '手写模板' and template.data['optimization']['mode'] == 'single'
    assert template.parameters[1]['values'] == ['甲', '乙']
    assert template.graph.evaluate({'quality': 3})['double_quality'].value == 6


def test_excel_errors_include_sheet_and_row_and_drafts_can_be_incomplete(tmp_path):
    data = blank_template(); data['parameters'] = [{'name': 'speed', 'unit': 'mm'}]
    file = tmp_path / '草稿.xlsx'; write_template_excel(data, file)
    assert read_template_excel(file)['parameters'][0]['name'] == 'speed'
    with pytest.raises(ValidationError): Template.read(file)
    workbook = load_workbook(file); workbook['输入参数']['F2'] = '错误范围'; workbook.save(file); workbook.close()
    with pytest.raises(ValidationError, match='输入参数 第 2 行'): read_template_excel(file)


def test_excel_roundtrip_does_not_change_implicit_model_roles(tmp_path):
    data = small_template(False).data
    data['measurements'].append({'name': 'other', 'unit': 'kg'})
    data['derived_metrics'].append({'name': 'other_twice', 'unit': 'kg', 'expression': 'other * 2'})
    data['measurements'][0].pop('model_base')
    original = Template(data); file = tmp_path / 'roles.xlsx'; original.write(file)
    restored = Template.read(file)
    assert [f['name'] for f in restored.responses] == [f['name'] for f in original.responses]
    assert not restored.measurements[-1]['is_response'] and not restored.metrics[-1]['is_response']


def test_cross_section_copy_leaves_only_common_properties_and_does_not_mutate():
    template = small_template().data; original = copy.deepcopy(template)
    for section in ('parameters', 'derived_metrics'):
        result, pasted = paste_fields(template, 'measurements', template['measurements'], section)
        field = pasted[0]
        assert field['unit'] == 'kg' and field['name'] == 'quality_copy1'
        assert not {'bounds', 'values', 'fixed_value', 'measured_at', 'model_base', 'is_response'} & field.keys()
        if section == 'derived_metrics': assert field['expression'] == ''
    assert template == original


def test_same_section_copy_remaps_selected_formula_dependencies():
    data = small_template().data
    data['derived_metrics'].append({'name': 'quadruple', 'expression': 'double_quality * 2', 'unit': 'kg'})
    result, pasted = paste_fields(data, 'derived_metrics', data['derived_metrics'], 'derived_metrics')
    assert pasted[1]['expression'] == 'double_quality_copy1 * 2'
    assert len(result['derived_metrics']) == 4


def test_free_text_records_predict_candidates_and_ax_tracking(tmp_path):
    template = small_template(); project = Project.create(tmp_path / 'text.sqlite', '文字实验', template)
    try:
        service = PlannerService(project)
        for speed in (1, 3, 6): service.add_record({'speed': speed, 'sample': '任意文字'}, {'quality': 2 * speed + 1})
        assert Encoder(template).width == 1
        model, encoder, names, _, _ = fit_models(template, project.experiments(), 'bayesian_linear_v1')
        result = prediction_summary(model, encoder, names, template, [{'speed': 4, 'sample': '另一个文字'}])
        assert result[0]['quality']['mean'] == pytest.approx(9, abs=.1)
        with pytest.raises(ValidationError, match='自由文字'): candidate_pool(template, BatchRequest())
        pool = candidate_pool(template, BatchRequest(baseline={'speed': 1, 'sample': '任意文字'}))
        assert pool and all(row['sample'] == '任意文字' for row in pool)
        experiment, mapping = rebuild(template, project.experiments())
        assert 'sample' not in experiment.search_space.parameters and len(mapping) == 3
        template.write(tmp_path / 'text.xlsx')
        assert Template.read(tmp_path / 'text.xlsx').parameters[1]['value_type'] == 'text'
    finally: project.close()


def test_session_save_discard_new_file_overwrite_and_reopen(tmp_path):
    file = tmp_path / '同名.sqlite'
    with Project.create(file, '旧项目', small_template(False)) as old: before = old.snapshot()
    with ProjectSession(file) as session:
        service = PlannerService(session); service.add_record({'speed': 1}, {'quality': 3})
        assert session.dirty
        with Project(file) as disk: assert disk.snapshot() == before
        session.save(); assert not session.dirty
        service.add_record({'speed': 2}, {'quality': 5})
    with Project(file) as disk: assert len(disk.experiments()) == 1
    with ProjectSession(file, creating=True, template=small_template(False)) as replacement:
        PlannerService(replacement).add_record({'speed': 3}, {'quality': 7})
        with Project(file) as disk: assert disk.name == '旧项目'
        replacement.save()
    with Project(file) as disk:
        assert disk.name == '同名' and disk.experiments()[0]['actual']['speed'] == 3
    new_file = tmp_path / '未保存.sqlite'
    with ProjectSession(new_file, creating=True): pass
    assert not new_file.exists()


def test_session_refuses_to_overwrite_external_changes(tmp_path):
    file = tmp_path / 'shared.sqlite'
    with Project.create(file, '共享', small_template(False)): pass
    with ProjectSession(file) as session:
        PlannerService(session).add_record({'speed': 1}, {'quality': 3})
        with Project(file) as external: PlannerService(external).add_record({'speed': 2}, {'quality': 5})
        with pytest.raises(StaleVersionError): session.save()
    with Project(file) as disk: assert disk.experiments()[0]['actual']['speed'] == 2


def test_failed_atomic_save_keeps_original_and_dirty_session(tmp_path, monkeypatch):
    import experiment_planner.storage.session as session_module
    file = tmp_path / 'fault.sqlite'
    with Project.create(file, '故障测试', small_template(False)): pass
    before = file.read_bytes()
    with ProjectSession(file) as session:
        PlannerService(session).add_record({'speed': 1}, {'quality': 3})
        def fail(*args): raise OSError('模拟保存失败')
        monkeypatch.setattr(session_module.os, 'replace', fail)
        with pytest.raises(OSError): session.save()
        assert session.dirty and file.read_bytes() == before
        assert len(session.experiments()) == 1
    assert not list(tmp_path.glob('.project-save-*'))
