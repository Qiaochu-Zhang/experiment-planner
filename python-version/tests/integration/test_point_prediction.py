"""Batch predictions keep exact conditions, source identity and cached Excel values."""
import copy

import pytest
from openpyxl import load_workbook

from experiment_planner.application.analysis import analyze
from experiment_planner.domain.errors import ValidationError
from experiment_planner.domain.template import Template
from experiment_planner.io.prediction_export import export_point_predictions


def point_template(uncertainty=False):
    return Template({"schema_version": 2, "template_version": 1, "name": "合成点预测",
        "parameters": [{"name": "x", "label": "流量", "unit": "sccm", "bounds": [0., 1.],
                        "execution_rounding": {"mode": "digits", "digits": 1}, "display_digits": 0,
                        "input_uncertainty": {"enabled": True, "kind": "std", "amount": .02}}],
        "measurements": [{"name": "y", "label": "结果", "unit": "nm", "is_response": True}],
        "derived_metrics": [{"name": "twice_y", "label": "双倍结果", "unit": "nm", "expression": "2*y"}],
        "objectives": [{"metric": "y", "direction": "maximize"}],
        "optimization": {"mode": "single"}, "ratio_policy": {"mode": "remove"},
        "input_uncertainty": {"enabled": uncertainty, "samples": 8}})


def point_snapshot(template):
    return {"name": "合成批量预测", "template": template.data, "revision": 8,
            "experiments": [{"id": i, "record_uid": f"uid-{i}", "status": "completed",
                "actual": {"x": x}, "derived": {}, "observations": {"y": {"value": 1+2*x+i*.01,
                    "uncertainty": {"kind": "std", "amount": .03}}}}
                for i, x in enumerate([.123456789123, .4, .8, .123456789123], 1)]}


def sheet_rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


@pytest.mark.model
@pytest.mark.parametrize("model", ["gp_rbf_v1", "gp_matern25_v1", "bayesian_linear_v1"])
@pytest.mark.parametrize("uncertainty", [False, True])
def test_fit_once_predict_all_exact_conditions_and_export_cached_values(tmp_path, monkeypatch, model, uncertainty):
    from experiment_planner.engine import models
    from experiment_planner.engine.objectives import prediction_summary
    from experiment_planner.plots import trends

    template = point_template(uncertainty)
    snapshot = point_snapshot(template)
    selected = [snapshot["experiments"][i] for i in (0, 1, 3)]
    conditions = [r["actual"] for r in selected]
    sources = [{"experiment_id": r["id"], "record_uid": r["record_uid"], "copied_revision": 8} for r in selected]
    request = {"conditions": conditions, "source_experiments": sources, "model": model, "seed": 3}
    original = copy.deepcopy((snapshot, request))
    fitted = []
    fit = models.fit_models
    def fit_once(*args, **kwargs):
        result = fit(*args, **kwargs)
        fitted.append(result)
        return result
    monkeypatch.setattr(models, "fit_models", fit_once)
    def no_plot(*args, **kwargs):
        pytest.fail("Batch point prediction must not generate a plot")
    monkeypatch.setattr(trends, "plot_slice", no_plot)
    report = analyze(snapshot, request)
    assert len(fitted) == 1
    assert (snapshot, request) == original
    assert report["conditions"] == conditions
    assert report["source_experiments"] == sources
    assert report["source_experiments"] is not sources
    assert "plot" not in report and "plot_path" not in report
    fitted_model, encoder, names, _, _ = fitted[0]
    assert report["predictions"] == prediction_summary(fitted_model, encoder, names, template, conditions, seed=3)
    path = tmp_path / "批量预测.xlsx"
    export_point_predictions(report, template, path)
    assert len(fitted) == 1  # Export never refits.
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["实验点预测", "预测明细", "字段说明", "预测说明"]
        rows = sheet_rows(workbook.active)
        assert [r["来源实验编号（复制时）"] for r in rows] == [1, 2, 4]
        assert [r["来源记录唯一标识"] for r in rows] == ["uid-1", "uid-2", "uid-4"]
        for row, condition, prediction in zip(rows, conditions, report["predictions"]):
            key = next(k for k in row if k.startswith("条件："))
            assert row[key] == condition["x"]
            for stat, expected in (("预测均值", prediction["y"]["mean"]),
                                   ("下限2.5%", prediction["y"]["quantiles"][0]),
                                   ("上限97.5%", prediction["y"]["quantiles"][2])):
                key = next(k for k in row if "[y]" in k and k.endswith(stat))
                assert row[key] == pytest.approx(expected, rel=1e-14)
        details = sheet_rows(workbook["预测明细"])
        assert {r["计算口径"] for r in details} == ({"输入确定", "含输入执行波动"} if uncertainty else {"输入确定"})
        assert {r["指标"] for r in details} == {"y", "twice_y"}
        assert any(r["结果类别"] == "目标1" for r in details)
        metadata = {r["项目"]: r["内容"] for r in sheet_rows(workbook["预测说明"])}
        assert metadata["预测时数据版本"] == 8
        assert metadata["模型"] == model and metadata["随机种子"] == 3
        if uncertainty: assert "点1输入波动配置与采样" in metadata
    finally:
        workbook.close()


@pytest.mark.parametrize("sources", [[], [{"experiment_id": 0}], [{"experiment_id": True}],
                                     [{"experiment_id": "1"}], [None]])
def test_bad_source_mapping_rejected_before_fit(monkeypatch, sources):
    from experiment_planner.engine import models
    monkeypatch.setattr(models, "fit_models", lambda *a, **k: pytest.fail("Must validate before fitting"))
    with pytest.raises(ValidationError, match="来源实验"):
        analyze(point_snapshot(point_template()), {"conditions": [{"x": .5}], "source_experiments": sources})


def test_invalid_condition_identifies_row_before_fit(monkeypatch):
    from experiment_planner.engine import models
    monkeypatch.setattr(models, "fit_models", lambda *a, **k: pytest.fail("Must validate before fitting"))
    with pytest.raises(ValidationError, match="预测点 2"):
        analyze(point_snapshot(point_template()), {"conditions": [{"x": .5}, {"x": 2.}]})


def test_mixed_condition_types_full_precision_and_literal_excel_text(tmp_path):
    template = point_template().data
    template["parameters"] += [{"name": "fixed", "fixed_value": 5.0},
        {"name": "gas", "value_type": "category", "values": ["A", "B"]},
        {"name": "enabled", "value_type": "bool"},
        {"name": "note", "value_type": "text", "use_as_model_input": False}]
    template = Template(template)
    condition = {"x": .123456789123, "fixed": 5., "gas": "B", "enabled": False, "note": "=1+1"}
    stats = {"mean": 9.87654321234, "quantiles": [1.123456789, 9.123456789, 15.7654321]}
    report = {"revision": 1, "template_version": 1, "model": "gp_rbf_v1", "conditions": [condition],
              "predictions": [{"y": stats}], "notices": ["=不要执行公式"]}
    path = tmp_path / "混合.xlsx"
    export_point_predictions(report, template, path)
    workbook = load_workbook(path)
    try:
        [row] = sheet_rows(workbook.active)
        for name, value in condition.items():
            key = next(k for k in row if k.startswith("条件：") and f"[{name}]" in k)
            assert row[key] == value
        assert all(c.data_type != "f" for sheet in workbook for line in sheet for c in line)
        assert all(sheet.freeze_panes == "A2" for sheet in workbook)
    finally:
        workbook.close()


@pytest.mark.parametrize("unavailable", [False, True])
def test_missing_ratio_mean_and_invalid_intervals_stay_blank(tmp_path, unavailable):
    stats = {"mean": None, "quantiles": None if unavailable else [1., 2., 3.],
             "nonfinite_sample_count": 4 if unavailable else 1,
             "moment_notice": "原始比值不声明均值存在", "summary_notice": "区间不可用" if unavailable else "仅有限样本"}
    report = {"revision": 1, "template_version": 1, "model": "gp_rbf_v1", "conditions": [{"x": .5}],
              "predictions": [{"raw_selectivity": stats}]}
    path = tmp_path / "比值.xlsx"
    export_point_predictions(report, point_template(), path)
    workbook = load_workbook(path)
    try:
        [row] = sheet_rows(workbook["预测明细"])
        assert row["预测均值"] is None
        assert row["下限2.5%"] == (None if unavailable else 1.)
        assert row["预测状态"] == ("区间不可用" if unavailable else "仅有限样本")
        assert "不声明均值存在" in row["说明"]
    finally:
        workbook.close()


@pytest.mark.parametrize("report", [{}, {"conditions": [{"x": .5}], "predictions": []},
    {"conditions": [{"x": .5}], "predictions": [{}]},
    {"conditions": [{"x": .5}], "predictions": [{"y": {"quantiles": [1, 2, 3]}}], "source_experiments": []}])
def test_incomplete_results_do_not_create_file(tmp_path, report):
    path = tmp_path / "不存在.xlsx"
    with pytest.raises(ValidationError): export_point_predictions(report, point_template(), path)
    assert not path.exists()


def test_only_excel_export_is_accepted(tmp_path):
    with pytest.raises(ValidationError, match="Excel"):
        export_point_predictions({}, point_template(), tmp_path / "结果.csv")
    assert not (tmp_path / "结果.csv").exists()
